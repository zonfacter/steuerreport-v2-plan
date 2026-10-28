import base64
import csv
import importlib.util
from io import BytesIO, StringIO
import json
from pathlib import Path
import re

from openpyxl import load_workbook
import pytest

from steuerreport_v2.db.duckdb_store import DuckDbStore
from steuerreport_v2.importers.csv_raw_importer import CsvRawImporter
from steuerreport_v2.importers.file_raw_importer import FileRawImporter
from steuerreport_v2.importers.raw_normalizer import RawNormalizer
from steuerreport_v2.importers.source_profiles import PROFILES, inspect_profile_columns


ROOT = Path(__file__).resolve().parents[2]
SPEC = importlib.util.spec_from_file_location("import_format_docs", ROOT / "scripts/generate_import_format_docs_v2.py")
DOCS = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(DOCS)


def test_documentation_and_templates_match_actual_profiles():
    for name, content in DOCS.render().items():
        assert (ROOT / name).read_bytes() == content, name
    wizard = (ROOT / "frontend/src/FileImportWizard.tsx").read_text()
    listed = re.search(r"const SOURCES = new Set\((\[.*?\])\);", wizard).group(1)
    assert set(DOCS.SOURCES) == set(json.loads(listed))
    assert "blockpit" not in DOCS.SOURCES
    assert (ROOT / "frontend/public/import-vorlage.csv").read_text() == DOCS.csv_text("manual-csv")


def test_downloaded_workbook_is_reproducible_and_contains_only_headers():
    first = DOCS.workbook_bytes()
    assert first == DOCS.workbook_bytes()
    book = load_workbook(BytesIO(first), read_only=True, data_only=False)
    assert book.sheetnames == list(DOCS.SOURCES)
    for source in DOCS.SOURCES:
        rows = list(book[source].values)
        assert rows == [tuple(next(csv.reader([DOCS.LAYOUTS[source]["headers"]])))]


@pytest.mark.parametrize("source", DOCS.SOURCES)
def test_csv_template_has_only_recognized_headers_and_no_transactions(source):
    rows = list(csv.reader(StringIO(DOCS.csv_text(source))))
    assert len(rows) == 1
    _, missing = inspect_profile_columns(PROFILES[source], rows[0])
    assert missing == []


@pytest.mark.parametrize("source", DOCS.SOURCES)
@pytest.mark.parametrize("file_format", ("csv", "xlsx"))
def test_documented_examples_preview_and_normalize_with_real_importer(tmp_path, source, file_format):
    store = DuckDbStore(tmp_path / "documented.duckdb")
    importer = FileRawImporter(CsvRawImporter(store))
    sample = DOCS.csv_text(source, example=True)
    if file_format == "csv":
        content = sample.encode("utf-8-sig")
    else:
        book = load_workbook(BytesIO(DOCS.workbook_bytes()))
        book[source].append(next(csv.reader([DOCS.LAYOUTS[source]["example"]])))
        buffer = BytesIO()
        book.save(buffer)
        content = buffer.getvalue()
    args = dict(source=source, filename=f"beispiel.{file_format}", content_base64=base64.b64encode(content).decode())
    if file_format == "xlsx":
        args["sheet_name"] = source
    if source == "mexc":
        args["mexc"] = {"account": "synthetic-spot", "export_kind": "spot_statement"}
    preview = importer.preview_base64(**args)
    assert preview.can_import
    assert preview.row_count == 1
    assert preview.missing_required_groups == []
    assert store.list_import_batches() == []
    batch = importer.import_base64(**args)
    events, report = RawNormalizer(store).normalize_batch_with_report(batch.batch_id)
    assert report.skipped_raw_objects == 0
    assert len(events) == 1
    event = events[0]
    assert event.event_time.isoformat() == "2025-01-15T10:00:00+00:00"
    if source in {"manual-csv", "coinbase"}:
        assert event.event_type == "trade"
        assert str(event.quantity) == "0.001"
        assert event.quote_asset == "EUR"
        assert event.quote_quantity == 70
        assert event.fee_asset == "EUR"
        assert str(event.fee_quantity) == "0.10"
    elif source == "helium":
        assert event.event_type == "mining"
        assert str(event.quantity) == "0.25"
    elif source == "mexc":
        assert event.event_type == "unknown"
        assert event.quantity == 2.5
    else:
        assert event.event_type == "transfer"
        if source == "solana_rpc":
            assert str(event.quantity) == "-0.1"
        else:
            assert event.quantity == 100


@pytest.mark.parametrize("source", DOCS.SOURCES)
def test_empty_template_is_not_importable(tmp_path, source):
    store = DuckDbStore(tmp_path / "empty.duckdb")
    importer = FileRawImporter(CsvRawImporter(store))
    preview = importer.preview_base64(
        source=source, filename="vorlage.csv",
        content_base64=base64.b64encode(DOCS.csv_text(source).encode()).decode(),
    )
    assert not preview.can_import
    assert preview.row_count == 0
    assert store.list_import_batches() == []


def test_type_is_not_a_technical_requirement_but_missing_fields_are():
    profile = PROFILES["manual-csv"]
    assert inspect_profile_columns(profile, ["time", "asset", "amount"])[1] == []
    for removed, required in (("time", "time"), ("asset", "asset"), ("amount", "quantity")):
        assert inspect_profile_columns(profile, [c for c in ("time", "asset", "amount") if c != removed])[1] == [required]


def test_shipped_documentation_pdfs_exist():
    for name in ("benutzerleitfaden.pdf", "importformate.pdf"):
        assert (ROOT / "frontend/public" / name).read_bytes().startswith(b"%PDF-")
