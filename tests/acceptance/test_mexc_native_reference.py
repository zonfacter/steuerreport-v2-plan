"""Unabhaengige, rein synthetische XLSX-Abnahme fuer GitHub Issue #4.

Sollwerte sind vor Backend-Implementierung festgelegt. Kein Skip/Xfail fuer
fehlende MEXC-Unterstuetzung; keine privaten Exporte, Preise oder Konten.
"""

import base64
from concurrent.futures import ThreadPoolExecutor
import csv
from datetime import datetime, timezone
from decimal import Decimal
import hashlib
from io import BytesIO, StringIO
import json
import os
from pathlib import Path
import re
import socket
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED

import duckdb
from openpyxl import Workbook
from pydantic import ValidationError
import pytest

from steuerreport_v2.contracts.imports import ImportBatchCreate, RawObjectCreate
from steuerreport_v2.db.duckdb_store import DuckDbStore
from steuerreport_v2.importers.csv_raw_importer import CsvRawImporter
from steuerreport_v2.importers.file_raw_importer import FileRawImporter
from steuerreport_v2.importers.raw_normalizer import RawNormalizer
from steuerreport_v2.ledger.balance import BalanceLedgerBuilder


D = Decimal
REPO = Path(__file__).resolve().parents[2]
UID = "SYNTHETIC-ACCOUNT"
ACCOUNT = "synthetic-spot"
LEDGER_ACCOUNT = "mexc:" + UID + ":" + ACCOUNT
WHEN = datetime(2031, 7, 9, 11, tzinfo=timezone.utc)
REASON = "Synthetische Originalbelege einzeln geprueft."
SPOT_HEADERS = (
    "UID", "Creation Time(UTC+01:00)", "Crypto", "Transaction Type",
    "Direction", "Quantity",
)
HEADERS = {
    "spot_statement": SPOT_HEADERS,
    "funding_other": ("UID", "Time", "Crypto", "Type", "Quantity", "Status", "Remark"),
    "asset_snapshot": ("UID", "Date(UTC+8)", "Account", "Crypto", "Total", "USDT Valuation"),
    "futures_capital_flow": (
        "UID", "Time(UTC+01:00)", "Futures Trading Pair", "Crypto", "Fund Type",
        "Fund Flow Type", "Amount",
    ),
    **{kind: SPOT_HEADERS for kind in (
        "futures_statement", "earn_fixed", "earn_flexible", "send_history",
    )},
}


@pytest.fixture
def case(tmp_path, monkeypatch):
    assert not tmp_path.resolve().is_relative_to(REPO)
    profile = tmp_path / "profile.json"
    profile.write_text("{}\n", encoding="ascii")
    monkeypatch.setenv("STEUERREPORT_PRIVATE_PROFILE", str(profile))
    monkeypatch.setenv("STEUERREPORT_PRIVATE_CONFIG_DIR", str(tmp_path / "configs"))
    monkeypatch.setenv("STEUERREPORT_STATE_DIR", str(tmp_path / "runtime"))
    monkeypatch.setenv("STEUERREPORT_V2_DB", str(tmp_path / "synthetic.duckdb"))
    original_init = DuckDbStore.__init__
    original_connect = duckdb.connect
    original_socket_connect = socket.socket.connect
    original_socket_connect_ex = socket.socket.connect_ex

    def check_path(path):
        assert Path(path).resolve().is_relative_to(tmp_path.resolve()), "Fremde DB gesperrt"

    def guarded_init(self, path=None):
        if path is not None:
            check_path(path)
        original_init(self, path)
        check_path(self.path)

    def guarded_connect(database=":memory:", *args, **kwargs):
        if str(database) != ":memory:":
            check_path(database)
        return original_connect(database, *args, **kwargs)

    def blocked(*args, **kwargs):
        raise OSError("MEXC-Referenz: kein Netzwerk/DNS")

    def connect(sock, address):
        if sock.family in {socket.AF_INET, socket.AF_INET6}:
            return blocked()
        return original_socket_connect(sock, address)

    def connect_ex(sock, address):
        if sock.family in {socket.AF_INET, socket.AF_INET6}:
            return blocked()
        return original_socket_connect_ex(sock, address)

    monkeypatch.setattr(DuckDbStore, "__init__", guarded_init)
    monkeypatch.setattr(duckdb, "connect", guarded_connect)
    monkeypatch.setattr(socket.socket, "connect", connect)
    monkeypatch.setattr(socket.socket, "connect_ex", connect_ex)
    monkeypatch.setattr(socket, "getaddrinfo", blocked)
    store = DuckDbStore(tmp_path / "synthetic.duckdb")
    store.health()
    return store, FileRawImporter(CsvRawImporter(store))


def _options(kind="spot_statement", **changes):
    result = {"export_kind": kind, "account": ACCOUNT}
    if kind == "funding_other":
        result["timezone"] = "UTC+01:00"
    result.update(changes)
    return result


def _spot(**changes):
    result = dict(zip(SPOT_HEADERS, (
        UID, "2031-07-09 12:00:00", "DEMOCOIN", "Airdrop", "Inflow", "2.5",
    )))
    result.update(changes)
    return result


def _funding(**changes):
    result = dict(zip(HEADERS["funding_other"], (
        UID, "2031-07-09 12:00:00", "DEMOCOIN", "Airdrop", "2.5", None, None,
    )))
    result.update(changes)
    return result


def _snapshot_row():
    return dict(zip(HEADERS["asset_snapshot"], (UID, "20310709", "Spot", "DEMOCOIN", "2.5", "123")))


def _bonus_rows():
    return [dict(zip(HEADERS["futures_capital_flow"], values)) for values in (
        (UID, "2031-07-09 12:00:00", None, "USDT", "BONUS", "Inflow", "7"),
        (UID, "2031-07-10 12:00:00", None, "USDT", "BONUS", "Outflow", "-7"),
    )]


def _xlsx(rows, kind="spot_statement", *, sheets=("Sheet1",), revision="original", extra_headers=()):
    book = Workbook()
    book.properties.creator = "SYNTHETIC TEST ONLY"
    book.properties.title = revision
    book.properties.created = datetime(2031, 7, 9)
    book.properties.modified = datetime(2031, 7, 9)
    for index, name in enumerate(sheets):
        sheet = book.active if index == 0 else book.create_sheet()
        sheet.title = name
        headers = HEADERS[kind] + tuple(extra_headers)
        sheet.append(headers)
        for row in rows:
            sheet.append([row.get(header) for header in headers])
    buffer = BytesIO()
    book.save(buffer)
    book.close()
    return buffer.getvalue()


def _multi_sheet_xlsx(sheets):
    book = Workbook()
    for index, (name, kind, rows) in enumerate(sheets):
        sheet = book.active if index == 0 else book.create_sheet()
        sheet.title = name
        sheet.append(HEADERS[kind])
        for row in rows:
            sheet.append([row.get(header) for header in HEADERS[kind]])
    buffer = BytesIO()
    book.save(buffer)
    book.close()
    return buffer.getvalue()


def _numeric_xlsx(uid, quantity):
    source = _xlsx([_spot()])
    output = BytesIO()
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    with ZipFile(BytesIO(source)) as original, ZipFile(output, "w") as changed:
        for member in original.infolist():
            content = original.read(member.filename)
            if member.filename == "xl/worksheets/sheet1.xml":
                root = ET.fromstring(content)
                for coordinate, literal in (("A2", uid), ("F2", quantity)):
                    cell = root.find(f".//{{{namespace}}}c[@r='{coordinate}']")
                    assert cell is not None
                    cell.clear()
                    cell.attrib.update(r=coordinate, t="n")
                    ET.SubElement(cell, "{" + namespace + "}v").text = literal
                content = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            changed.writestr(member, content)
    return output.getvalue()


def _csv(rows, kind="funding_other"):
    output = StringIO(newline="")
    writer = csv.DictWriter(output, fieldnames=HEADERS[kind], lineterminator="\r\n")
    writer.writeheader()
    writer.writerows(rows)
    return output.getvalue().encode("utf-8")


def _request(content, options, *, filename="synthetic.xlsx", sheet="Sheet1"):
    return {
        "source": "mexc", "filename": filename,
        "content_base64": base64.b64encode(content).decode("ascii"),
        "sheet_name": sheet, "mexc": options,
    }


def _bind_decisions(case, content, options, **kwargs):
    if options is None:
        return options
    decisions = options.get("decisions", {}) if isinstance(options, dict) else options.decisions
    missing = [
        (number, decision) for number, decision in decisions.items()
        if ("input_digest" not in decision if isinstance(decision, dict)
            else "input_digest" not in decision.model_fields_set)
    ]
    if not missing:
        return options
    without_decisions = (
        {key: value for key, value in options.items() if key != "decisions"}
        if isinstance(options, dict) else options.model_copy(update={"decisions": {}})
    )
    preview = case[1].preview_base64(**_request(content, without_decisions, **kwargs))
    rows = _native(preview)["rows"]
    # Exakte Wiederholung wendet keine neue Entscheidung an und bleibt unveraendert.
    if preview.duplicate_batch_id is not None and all(not row["input_digest"] and not row["candidates"] for row in rows):
        return options
    preview_rows = {str(row["row_number"]): row for row in rows}
    for number, decision in missing:
        if str(number) not in preview_rows:
            continue
        row = preview_rows[str(number)]
        digest = row["input_digest"]
        # Vorab ungueltige Zeilen besitzen noch keinen bindbaren Kandidatenstand.
        if digest == "":
            assert row["status"] in {"invalid", "review", "evidence_only"}
            assert not row["candidates"]
            continue
        assert isinstance(digest, str) and re.fullmatch(r"[0-9a-fA-F]{64}", digest)
        if isinstance(decision, dict):
            decision["input_digest"] = digest
        else:
            decision.input_digest = digest
    return options


def _preview(case, content, options, *, bind_decisions=True, **kwargs):
    if bind_decisions:
        options = _bind_decisions(case, content, options, **kwargs)
    return case[1].preview_base64(**_request(content, options, **kwargs))


def _native(preview):
    native = preview.model_dump(mode="json")["mexc"]
    assert isinstance(native, dict)
    assert {"raw_rows", "bookable_rows", "linked_rows", "review_rows", "snapshot_rows", "bonus_rows", "rows"} <= native.keys()
    return native


def _counts(native, **expected):
    for name, value in expected.items():
        assert native[name] == value, (name, native)


def _import(case, content, options, *, bind_decisions=True, **kwargs):
    store, importer = case
    if bind_decisions:
        options = _bind_decisions(case, content, options, **kwargs)
    batch = importer.import_base64(**_request(content, options, **kwargs))
    events = RawNormalizer(store).normalize_batch(batch.batch_id)
    BalanceLedgerBuilder(store).rebuild()
    return batch, events


def _raw(case, batch):
    return case[0].list_raw_objects(batch.batch_id, limit=None)


def _metadata(raw):
    payload = json.loads(raw.payload_json)
    assert isinstance(payload["_mexc"], str), "_mexc muss serialisiertes JSON sein"
    metadata = json.loads(payload["_mexc"])
    assert metadata["version"] == 1
    assert metadata["row_number"] == raw.row_number
    return payload, metadata


def _positions(case):
    return {(p.account, p.asset): p.quantity for p in BalanceLedgerBuilder(case[0]).positions()}


def _state(case):
    with case[0].connect() as con:
        return {table: con.execute(f"select * from {table} order by all").fetchall() for table in (
            "raw_import_batches", "raw_objects", "canonical_events", "balance_ledger",
            "review_decisions", "tax_lots", "tax_lines", "derivative_lines",
        )}


def _reject_before_batch(case, monkeypatch, content, options, *, bind_decisions=True, **kwargs):
    before = _state(case)

    def forbidden(*args, **kwargs):
        pytest.fail("Ungueltiger Import darf nicht einmal einen Batch anlegen")

    with monkeypatch.context() as patch:
        patch.setattr(case[0], "create_import_batch", forbidden)
        if hasattr(case[0], "create_import_batch_with_raw"):
            patch.setattr(case[0], "create_import_batch_with_raw", forbidden)
        with pytest.raises(ValueError):
            if bind_decisions:
                options = _bind_decisions(case, content, options, **kwargs)
            case[1].import_base64(**_request(content, options, **kwargs))
    assert _state(case) == before


def _decision(action, target=None, reason=REASON):
    result = {"action": action, "reason": reason}
    if target is not None:
        result["target_raw_object_id"] = target
    return result


def _seed(case, *, row=None, options=None):
    batch, events = _import(case, _xlsx([row or _spot()]), options or _options())
    assert len(events) == 1
    return batch, _raw(case, batch)[0]


def test_harness_blocks_foreign_database_and_all_ip_network(case, tmp_path):
    with pytest.raises(AssertionError, match="Fremde DB"):
        DuckDbStore(tmp_path.parent / "forbidden.duckdb")
    with pytest.raises(AssertionError, match="Fremde DB"):
        duckdb.connect(str(tmp_path.parent / "forbidden.duckdb"))
    with pytest.raises(OSError, match="kein Netzwerk"):
        socket.getaddrinfo("synthetic.invalid", 443)
    with socket.socket() as sock:
        with pytest.raises(OSError, match="kein Netzwerk"):
            sock.connect(("127.0.0.1", 9))
        with pytest.raises(OSError, match="kein Netzwerk"):
            sock.connect_ex(("127.0.0.1", 9))


@pytest.mark.parametrize("kind", HEADERS)
def test_header_only_native_xlsx_is_empty_not_corrupt(case, kind):
    before = _state(case)
    preview = _preview(case, _xlsx([], kind), _options(kind))
    assert preview.file_format == "xlsx"
    assert preview.sheet_name == "Sheet1"
    assert preview.columns == list(HEADERS[kind])
    assert preview.row_count == 0
    assert preview.can_import is False
    native = _native(preview)
    _counts(native, raw_rows=0, bookable_rows=0, linked_rows=0, review_rows=0, snapshot_rows=0, bonus_rows=0)
    assert native["rows"] == []
    assert _state(case) == before


def test_multisheet_xlsx_requires_selection_then_previews_and_normalizes_selected_sheet(case, monkeypatch):
    content = _multi_sheet_xlsx((
        ("Spot", "spot_statement", [_spot()]),
        ("Funding", "funding_other", [_funding()]),
    ))
    before = _state(case)

    ambiguous = _preview(case, content, _options(), sheet=None)
    assert ambiguous.file_format == "xlsx"
    assert ambiguous.sheet_names == ["Spot", "Funding"]
    assert ambiguous.sheet_name is None
    assert ambiguous.row_count == 0 and ambiguous.sample_rows == []
    assert ambiguous.can_import is False
    assert _native(ambiguous)["detected_kind"] is None
    assert any("Tabellenblatt ausdruecklich auswaehlen" in reason for reason in _native(ambiguous)["blocking_reasons"])
    assert _state(case) == before
    _reject_before_batch(case, monkeypatch, content, _options(), sheet=None)

    options = _options("funding_other", timezone="UTC+01:00")
    selected = _preview(case, content, options, sheet="Funding")
    assert selected.can_import is True
    assert selected.sheet_name == "Funding"
    assert _native(selected)["detected_kind"] == "funding_other"
    assert selected.row_count == 1

    batch, events = _import(case, content, options, sheet="Funding")
    assert len(events) == 1
    assert (events[0].event_type, events[0].asset, events[0].quantity, events[0].event_time) == (
        "unknown", "DEMOCOIN", D("2.5"), WHEN,
    )
    raw, = _raw(case, batch)
    payload, meta = _metadata(raw)
    assert raw.row_number == 2
    assert meta["sheet_name"] == "Funding"
    assert payload["Type"] == "Airdrop"
    assert _positions(case) == {(LEDGER_ACCOUNT, "DEMOCOIN"): D("2.5")}
    assert _state(case)["tax_lots"] == []
    archived = case[0].path.parent / "mexc_originals" / (hashlib.sha256(content).hexdigest() + ".original")
    assert archived.read_bytes() == content


def test_missing_options_cannot_guess_product_from_header_or_filename(case, monkeypatch):
    content = _xlsx([_spot()])
    before = _state(case)
    preview = _preview(case, content, None, filename="MEXC Spot Statement.xlsx")
    assert preview.can_import is False
    assert _state(case) == before
    _reject_before_batch(case, monkeypatch, content, None)


@pytest.mark.parametrize("options", [{}, {"export_kind": "spot_statement"}, {"account": ACCOUNT}])
def test_incomplete_options_are_preview_blockers_not_silent_defaults(case, monkeypatch, options):
    content = _xlsx([_spot()])
    preview = _preview(case, content, options)
    assert preview.can_import is False
    assert _native(preview)["blocking_reasons"]
    _reject_before_batch(case, monkeypatch, content, options)


@pytest.mark.parametrize("changes", [
    {"export_kind": "unverified_product"},
    {"account": "synthetic\naccount"},
    {"timezone": "SYNTHETIC/NO_SUCH_ZONE"},
    {"decisions": {"2": {"action": "auto_merge", "reason": REASON}}},
    {"decisions": {"2": _decision("separate", reason="123456789")}},
])
def test_malformed_dict_options_raise_pydantic_validation_error(case, changes):
    before = _state(case)
    with pytest.raises(ValidationError):
        _preview(case, _xlsx([_spot()]), _options(**changes))
    assert _state(case) == before


def test_contract_model_and_dict_have_same_preview_and_assistant_preserves_options(case):
    from steuerreport_v2.contracts.imports import ImportAssistantRequest
    from steuerreport_v2.contracts.mexc_import import MexcImportOptions

    content = _xlsx([_spot()])
    options = _options()
    typed = MexcImportOptions.model_validate(options)
    request = ImportAssistantRequest(import_mode="file", **_request(content, options))
    assert isinstance(request.mexc, MexcImportOptions)
    assert request.mexc == typed
    assert _preview(case, content, typed) == _preview(case, content, options)
    batch, events = _import(case, content, typed)
    assert len(events) == 1
    assert _metadata(_raw(case, batch)[0])[1]["export_kind"] == "spot_statement"


def test_airdrop_ready_and_review_overlap_raw_provenance_and_fixed_summer_offset(case):
    row = _spot()
    content = _xlsx([row])
    before = _state(case)
    preview = _preview(case, content, _options())
    native = _native(preview)
    assert preview.can_import is True
    _counts(native, raw_rows=1, bookable_rows=1, linked_rows=0, review_rows=1, snapshot_rows=0, bonus_rows=0)
    assert native["rows"][0]["row_number"] == 2
    assert native["rows"][0]["status"] == "ready"
    assert _state(case) == before, "Vorschau muss rein lesend sein"
    batch, events = _import(case, content, _options())
    assert len(events) == 1
    event = events[0]
    assert (event.source, event.account, event.asset, event.event_type, event.quantity) == (
        "mexc", LEDGER_ACCOUNT, "DEMOCOIN", "unknown", D("2.5"),
    )
    assert event.event_time.astimezone(timezone.utc) == WHEN
    assert event.quote_asset is None and event.quote_quantity is None
    assert event.fee_quantity in (None, D("0"))
    raw, = _raw(case, batch)
    assert raw.row_number == 2
    assert event.raw_object_id == raw.raw_object_id
    payload, meta = _metadata(raw)
    assert {key: payload[key] for key in row} == row
    assert meta["export_kind"] == "spot_statement"
    assert meta["account"] == ACCOUNT
    assert meta["uid"] == UID
    assert meta["sheet_name"] == "Sheet1"
    assert meta["file_sha256"] == "sha256:" + hashlib.sha256(content).hexdigest()
    assert meta["file_sha256"] == preview.file_sha256
    assert meta["role"] == "book"
    assert datetime.fromisoformat(meta["event_time"].replace("Z", "+00:00")).astimezone(timezone.utc) == WHEN
    assert (meta["asset"], D(meta["quantity"]), meta["label"]) == ("DEMOCOIN", D("2.5"), "Airdrop")
    assert meta["signature"]
    assert not meta.get("link_target_raw_object_id")
    assert _positions(case) == {(LEDGER_ACCOUNT, "DEMOCOIN"): D("2.5")}
    archived = case[0].path.parent / "mexc_originals" / (hashlib.sha256(content).hexdigest() + ".original")
    assert archived.read_bytes() == content


@pytest.mark.parametrize("quantity,direction,expected", [
    ("2.5", "Inflow", "2.5"), ("+2.5", "Inflow", "2.5"),
    ("2.5", "Outflow", "-2.5"), ("-2.5", "Outflow", "-2.5"),
    ("2.500000000001", "Inflow", "2.500000000001"),
])
def test_direction_sign_and_decimal_values_are_not_double_negated_or_rounded(case, quantity, direction, expected):
    row = _spot(Quantity=quantity, Direction=direction)
    batch, events = _import(case, _xlsx([row]), _options())
    assert len(events) == 1
    assert events[0].quantity == D(expected)
    assert events[0].event_type == "unknown"
    assert _positions(case) == {(LEDGER_ACCOUNT, "DEMOCOIN"): D(expected)}
    assert _metadata(_raw(case, batch)[0])[0]["Quantity"] == quantity


@pytest.mark.parametrize("quantity,direction,expected", [
    ("12345678901234567890.123456789012345678", "Inflow", "12345678901234567890.123456789012345678"),
    ("12345678901234567890.123456789012345678", "Outflow", "-12345678901234567890.123456789012345678"),
    ("99999999999999999999.999999999999999999", "Inflow", "99999999999999999999.999999999999999999"),
])
def test_full_decimal_38_18_precision_survives_native_canonical_and_ledger(case, quantity, direction, expected):
    content = _xlsx([_spot(Quantity=quantity, Direction=direction)])
    batch, events = _import(case, content, _options())
    assert len(events) == 1
    assert events[0].quantity == D(expected), "Canonical darf nicht auf Decimal-Defaultpraezision runden"
    assert D(_metadata(_raw(case, batch)[0])[1]["quantity"]) == D(expected)
    assert _positions(case) == {(LEDGER_ACCOUNT, "DEMOCOIN"): D(expected)}
    with case[0].connect() as con:
        stored = con.execute("select delta_quantity,balance_quantity from balance_ledger").fetchall()
    assert stored == [(D(expected), D(expected))], "Auch die laufende Ledger-Balance muss exakt bleiben"


def test_numeric_excel_xml_literals_preserve_large_uid_and_eighteen_decimal_places(case):
    uid, quantity = "9007199254740993", "0.123456789012345678"
    content = _numeric_xlsx(uid, quantity)
    batch, events = _import(case, content, _options())
    payload, meta = _metadata(_raw(case, batch)[0])
    assert (payload["UID"], payload["Quantity"]) == (uid, quantity)
    assert meta["uid"] == uid
    assert events[0].quantity == D(quantity)
    assert _positions(case) == {("mexc:" + uid + ":" + ACCOUNT, "DEMOCOIN"): D(quantity)}


@pytest.mark.parametrize("quantity", ["100000000000000000000", "0.1234567890123456789"])
def test_unrepresentable_decimal_values_block_before_batch_instead_of_rounding(case, monkeypatch, quantity):
    content = _xlsx([_spot(Quantity=quantity)])
    preview = _preview(case, content, _options())
    assert preview.can_import is False
    assert _native(preview)["rows"][0]["status"] == "invalid"
    _reject_before_batch(case, monkeypatch, content, _options())


@pytest.mark.parametrize("limit", ["rows", "file_bytes", "expanded_zip"])
def test_native_size_limits_reject_without_partial_import(case, monkeypatch, limit):
    if limit == "rows":
        content = _xlsx([_spot()] * 10001)
    elif limit == "file_bytes":
        content = b"x" * (25 * 1024 * 1024 + 1)
    else:
        output = BytesIO()
        with ZipFile(BytesIO(_xlsx([_spot()]))) as original, ZipFile(output, "w", compression=ZIP_DEFLATED) as changed:
            for member in original.infolist():
                changed.writestr(member.filename, original.read(member.filename))
            # Nur 1 MiB gleichzeitig im RAM; hoch komprimierter synthetischer Ballast.
            with changed.open("synthetic-padding.bin", "w") as padding:
                chunk = b"0" * (1024 * 1024)
                for _ in range(101):
                    padding.write(chunk)
        content = output.getvalue()
        assert len(content) < 25 * 1024 * 1024
    _reject_before_batch(case, monkeypatch, content, _options())


@pytest.mark.parametrize("changes", [{"Quantity": "-2.5"}, {"Direction": "SIDEWAYS"}, {"Quantity": "not-a-number"}])
def test_contradictory_or_unknown_direction_and_bad_amount_are_not_bookable(case, monkeypatch, changes):
    content = _xlsx([_spot(**changes)])
    preview = _preview(case, content, _options())
    assert preview.can_import is False
    native = _native(preview)
    _counts(native, raw_rows=1, bookable_rows=0)
    row, = native["rows"]
    assert row["row_number"] == 2
    assert row["status"] in {"invalid", "review"}
    assert row["reason"]
    assert case[0].list_canonical_events() == []
    _reject_before_batch(case, monkeypatch, content, _options())


@pytest.mark.parametrize("missing", ["timezone", "time"])
def test_funding_without_time_basis_or_value_blocks_even_raw_import(case, monkeypatch, missing):
    options = _options("funding_other")
    row = _funding()
    if missing == "timezone":
        options.pop("timezone")
        _seed(case)
    else:
        row["Time"] = ""
    content = _xlsx([row], "funding_other")
    preview = _preview(case, content, options)
    assert preview.can_import is False
    native = _native(preview)
    _counts(native, raw_rows=1, bookable_rows=0)
    assert native["rows"][0]["reason"]
    _reject_before_batch(case, monkeypatch, content, options)


def test_funding_explicit_fixed_offset_books_unknown_not_income(case):
    batch, events = _import(case, _xlsx([_funding()], "funding_other"), _options("funding_other"))
    assert len(events) == 1
    assert events[0].event_time.astimezone(timezone.utc) == WHEN
    assert (events[0].event_type, events[0].quantity) == ("unknown", D("2.5"))
    payload, meta = _metadata(_raw(case, batch)[0])
    assert payload["Time"] == "2031-07-09 12:00:00"
    assert meta["timezone"] == "UTC+01:00"


def test_physical_row_numbers_survive_blank_excel_row(case):
    content = _xlsx([{}, _spot()])
    native = _native(_preview(case, content, _options()))
    _counts(native, raw_rows=1, bookable_rows=1)
    assert native["rows"][0]["row_number"] == 3
    batch, events = _import(case, content, _options())
    assert len(events) == 1
    raw, = _raw(case, batch)
    assert raw.row_number == 3
    assert _metadata(raw)[1]["row_number"] == 3


@pytest.mark.parametrize("local_time,expected_status", [("11:00:00", "duplicate_candidate"), ("12:00:00", "ready")])
def test_candidate_comparison_uses_utc_instant_not_local_time_text(case, local_time, expected_status):
    _seed(case)
    content = _xlsx([_funding(Time="2031-07-09 " + local_time)], "funding_other")
    native = _native(_preview(case, content, _options("funding_other", timezone="UTC")))
    assert native["rows"][0]["status"] == expected_status


def test_same_bytes_from_other_source_do_not_hide_mexc_file(case):
    content = _xlsx([_spot()])
    foreign = case[1].import_base64(source="binance", filename="synthetic.xlsx", content_base64=base64.b64encode(content).decode("ascii"))
    preview = _preview(case, content, _options())
    assert preview.duplicate_batch_id is None
    assert preview.can_import is True
    assert _native(preview)["rows"][0]["status"] == "ready"
    own, events = _import(case, content, _options())
    assert own.batch_id != foreign.batch_id
    assert len(events) == 1


def test_separate_decision_ten_character_boundary_is_accepted(case):
    _seed(case)
    options = _options("funding_other", decisions={"2": _decision("separate", reason="1234567890")})
    batch, events = _import(case, _xlsx([_funding()], "funding_other"), options)
    assert len(events) == 1
    assert _metadata(_raw(case, batch)[0])[1]["decision_reason"] == "1234567890"
    assert _positions(case) == {(LEDGER_ACCOUNT, "DEMOCOIN"): D("5.0")}


@pytest.mark.parametrize("path", ["base64", "csv_importer"])
def test_csv_native_path_preserves_physical_row_and_exact_original_bytes(case, path):
    out = StringIO(newline="")
    writer = csv.DictWriter(out, fieldnames=SPOT_HEADERS, lineterminator="\r\n")
    writer.writeheader()
    writer.writerow(_spot())
    text = out.getvalue()
    content = text.encode("utf-8")
    if path == "base64":
        batch, events = _import(case, content, _options(), filename="synthetic.csv", sheet=None)
    else:
        importer = CsvRawImporter(case[0])
        preview = importer.preview_text(source="mexc", filename="synthetic.csv", content=text, mexc=_options())
        assert preview.can_import is True
        batch = importer.import_text(source="mexc", filename="synthetic.csv", content=text, mexc=_options())
        events = RawNormalizer(case[0]).normalize_batch(batch.batch_id)
        BalanceLedgerBuilder(case[0]).rebuild()
    assert len(events) == 1
    assert events[0].quantity == D("2.5")
    raw, = _raw(case, batch)
    assert raw.row_number == 2
    assert _metadata(raw)[1]["sheet_name"] is None
    archived = case[0].path.parent / "mexc_originals" / (hashlib.sha256(content).hexdigest() + ".original")
    assert archived.read_bytes() == content


def test_snapshot_and_linked_raws_are_not_unmapped_but_bonus_is(case):
    from steuerreport_v2.domain.review import ReviewService

    snapshot, _ = _import(case, _xlsx([_snapshot_row()], "asset_snapshot"), _options("asset_snapshot"))
    _, root = _seed(case)
    linked_options = _options("funding_other", decisions={"2": _decision("link_existing", root.raw_object_id)})
    linked, _ = _import(case, _xlsx([_funding()], "funding_other"), linked_options)
    bonus, _ = _import(case, _xlsx(_bonus_rows(), "futures_capital_flow"), _options("futures_capital_flow"))
    unmapped = case[0].unmapped_raw_object_counts()
    assert snapshot.batch_id not in unmapped
    assert linked.batch_id not in unmapped
    assert unmapped[bonus.batch_id] == 2
    gate = ReviewService(case[0]).gate()
    assert gate.allow_export is False
    assert any(i.blocks_export and bonus.batch_id in i.scope for i in gate.issues)


@pytest.mark.parametrize("foreign_marker", ["bitget_withdrawal_evidence", "history_evidence_only"])
def test_mexc_bonus_cannot_escape_review_using_other_sources_evidence_markers(case, foreign_marker):
    from steuerreport_v2.domain.review import ReviewService

    rows = _bonus_rows()
    for row in rows:
        row["Futures Trading Pair"] = "SYNTHETIC_PAIR"
    extra = ()
    if foreign_marker == "bitget_withdrawal_evidence":
        for row in rows:
            row["Futures Trading Pair"] = foreign_marker
    else:
        extra = (foreign_marker,)
        for row in rows:
            row[foreign_marker] = "1"
    batch, events = _import(case, _xlsx(rows, "futures_capital_flow", extra_headers=extra), _options("futures_capital_flow"))
    assert events == []
    gate = ReviewService(case[0]).gate()
    observed = (case[0].unmapped_raw_object_counts().get(batch.batch_id), gate.allow_export)
    assert observed == (2, False), ("Fremde Belegmarker duerfen MEXC-BONUS nicht freigeben", observed, [i.code for i in gate.issues])
    assert any(i.blocks_export and batch.batch_id in i.scope for i in gate.issues)


def test_deleted_link_root_is_a_specific_blocker_and_never_recreated_as_booking(case):
    from steuerreport_v2.domain.review import ReviewService

    root_batch, root = _seed(case)
    linked_options = _options("funding_other", decisions={"2": _decision("link_existing", root.raw_object_id)})
    linked, _ = _import(case, _xlsx([_funding()], "funding_other"), linked_options)
    case[0].delete_import_batch(root_batch.batch_id)
    assert RawNormalizer(case[0]).normalize_batch(linked.batch_id) == []
    BalanceLedgerBuilder(case[0]).rebuild()
    assert _positions(case) == {}
    gate = ReviewService(case[0]).gate()
    assert gate.allow_export is False
    assert any(i.blocks_export and i.code == "mexc_link_target_missing" for i in gate.issues)


def test_cross_export_candidate_blocks_until_linked_then_keeps_two_raws_one_event(case, monkeypatch):
    root_batch, root = _seed(case)
    original = root.payload_json
    content = _xlsx([_funding()], "funding_other")
    options = _options("funding_other")
    preview = _preview(case, content, options)
    assert preview.can_import is False
    native = _native(preview)
    row, = native["rows"]
    assert row["status"] == "duplicate_candidate"
    assert row["row_number"] == 2
    assert row["candidates"]
    assert root.raw_object_id in set(_string_leaves(row["candidates"]))
    _reject_before_batch(case, monkeypatch, content, options)
    options["decisions"] = {"2": _decision("link_existing", root.raw_object_id)}
    linked_preview = _preview(case, content, options)
    assert linked_preview.can_import is True
    _counts(_native(linked_preview), raw_rows=1, bookable_rows=0, linked_rows=1)
    assert _native(linked_preview)["rows"][0]["status"] == "linked"
    linked_batch, events = _import(case, content, options)
    assert events == []
    assert len(case[0].list_canonical_events()) == 1
    linked_raw, = _raw(case, linked_batch)
    _, metadata = _metadata(linked_raw)
    assert metadata["role"] == "linked"
    assert metadata["link_target_raw_object_id"] == root.raw_object_id
    assert metadata["decision_reason"] == REASON
    assert metadata["signature"] == _metadata(root)[1]["signature"]
    assert _raw(case, root_batch)[0].payload_json == original
    assert sum(len(_raw(case, b)) for b in case[0].list_import_batches()) == 2
    assert _positions(case) == {(LEDGER_ACCOUNT, "DEMOCOIN"): D("2.5")}
    RawNormalizer(case[0]).normalize_batch(linked_batch.batch_id)
    RawNormalizer(case[0]).normalize_batch(root_batch.batch_id)
    BalanceLedgerBuilder(case[0]).rebuild()
    assert len(case[0].list_canonical_events()) == 1
    assert _positions(case) == {(LEDGER_ACCOUNT, "DEMOCOIN"): D("2.5")}


@pytest.mark.parametrize("action,expected_count,expected_quantity", [("link_existing", 1, "2.5"), ("separate", 2, "5.0")])
def test_csv_cross_file_decisions_are_checked_independently_of_xlsx_empty_cells(case, monkeypatch, action, expected_count, expected_quantity):
    root_batch, root = _seed(case)
    before_root = root.payload_json
    content = _csv([_funding()])
    options = _options("funding_other")
    kwargs = {"filename": "synthetic.csv", "sheet": None}
    preview = _preview(case, content, options, **kwargs)
    assert preview.can_import is False
    assert _native(preview)["rows"][0]["status"] == "duplicate_candidate"
    _reject_before_batch(case, monkeypatch, content, options, **kwargs)
    options["decisions"] = {"2": _decision(action, root.raw_object_id if action == "link_existing" else None)}
    batch, _ = _import(case, content, options, **kwargs)
    _, meta = _metadata(_raw(case, batch)[0])
    assert meta["role"] == ("linked" if action == "link_existing" else "book")
    assert meta["decision_reason"] == REASON
    if action == "link_existing":
        assert meta["link_target_raw_object_id"] == root.raw_object_id
    assert _raw(case, root_batch)[0].payload_json == before_root
    assert len(case[0].list_canonical_events()) == expected_count
    assert _positions(case) == {(LEDGER_ACCOUNT, "DEMOCOIN"): D(expected_quantity)}


def test_separate_decision_preserves_second_legitimate_equal_booking(case):
    _seed(case)
    options = _options("funding_other", decisions={"2": _decision("separate")})
    batch, events = _import(case, _xlsx([_funding()], "funding_other"), options)
    assert len(events) == 1
    assert len(case[0].list_canonical_events()) == 2
    assert _positions(case) == {(LEDGER_ACCOUNT, "DEMOCOIN"): D("5.0")}
    meta = _metadata(_raw(case, batch)[0])[1]
    assert meta["role"] == "book"
    assert meta["decision_reason"] == REASON


def test_two_identical_physical_rows_are_not_automatically_collapsed(case):
    content = _xlsx([_spot(), _spot()])
    batch, events = _import(case, content, _options())
    assert len(events) == 2
    assert len({e.raw_object_id for e in events}) == 2
    assert [r.row_number for r in _raw(case, batch)] == [2, 3]
    assert _positions(case) == {(LEDGER_ACCOUNT, "DEMOCOIN"): D("5.0")}


@pytest.mark.parametrize("action,expected_events,expected_balance", [("link_existing", 1, "2.5"), ("separate", 2, "5.0")])
def test_regenerated_file_different_hash_needs_explicit_overlap_decision(case, monkeypatch, action, expected_events, expected_balance):
    original = _xlsx([_spot()], revision="first")
    changed = _xlsx([_spot()], revision="regenerated")
    assert hashlib.sha256(original).digest() != hashlib.sha256(changed).digest()
    batch, _ = _import(case, original, _options())
    root, = _raw(case, batch)
    options = _options()
    preview = _preview(case, changed, options)
    assert preview.file_sha256 != _metadata(root)[1]["file_sha256"]
    assert preview.can_import is False
    assert _native(preview)["rows"][0]["status"] == "duplicate_candidate"
    _reject_before_batch(case, monkeypatch, changed, options)
    options["decisions"] = {"2": _decision(action, root.raw_object_id if action == "link_existing" else None)}
    _import(case, changed, options)
    assert len(case[0].list_canonical_events()) == expected_events
    assert _positions(case) == {(LEDGER_ACCOUNT, "DEMOCOIN"): D(expected_balance)}


def test_reimport_is_context_bound_idempotent_and_decisions_do_not_change_identity(case):
    content = _xlsx([_spot()])
    options = _options()
    fingerprint = _preview(case, content, options).import_fingerprint
    batch, _ = _import(case, content, options)
    before = _state(case)
    changed_decision = _options(decisions={"2": _decision("separate", reason="Anderer synthetischer Pruefvermerk.")})
    preview = _preview(case, content, changed_decision, filename="renamed.xlsx")
    assert preview.import_fingerprint == fingerprint
    assert preview.duplicate_batch_id == batch.batch_id
    repeated = case[1].import_base64(**_request(content, changed_decision, filename="renamed.xlsx"))
    assert repeated.batch_id == batch.batch_id
    assert _state(case) == before


def test_concurrent_identical_file_import_does_not_create_partial_or_double_batch(case):
    request = _request(_xlsx([_spot()]), _options())
    with ThreadPoolExecutor(max_workers=2) as pool:
        batches = list(pool.map(lambda _: case[1].import_base64(**request), range(2)))
    assert batches[0].batch_id == batches[1].batch_id
    assert len(case[0].list_import_batches()) == 1
    assert len(_raw(case, batches[0])) == 1
    events = RawNormalizer(case[0]).normalize_batch(batches[0].batch_id)
    BalanceLedgerBuilder(case[0]).rebuild()
    assert len(events) == 1
    assert _positions(case) == {(LEDGER_ACCOUNT, "DEMOCOIN"): D("2.5")}


def test_stale_ready_preview_is_rechecked_before_import(case, monkeypatch):
    later = _xlsx([_spot()], revision="later")
    assert _preview(case, later, _options()).can_import is True
    _seed(case)
    _reject_before_batch(case, monkeypatch, later, _options())


@pytest.mark.parametrize("dimension", ["account", "export_kind", "timezone", "sheet"])
def test_import_fingerprint_binds_each_mapping_dimension(case, dimension):
    kind = "funding_other" if dimension == "timezone" else "spot_statement"
    content = _xlsx([_funding() if kind == "funding_other" else _spot()], kind, sheets=("Sheet1", "Sheet2"))
    original_options = _options(kind)
    changed_options = dict(original_options)
    changed_sheet = "Sheet1"
    if dimension == "sheet":
        changed_sheet = "Sheet2"
    elif dimension == "account":
        changed_options["account"] = "synthetic-second-account"
    elif dimension == "export_kind":
        changed_options["export_kind"] = "futures_statement"
    else:
        changed_options["timezone"] = "UTC"
    original = _preview(case, content, original_options)
    changed = _preview(case, content, changed_options, sheet=changed_sheet)
    assert original.file_sha256 == changed.file_sha256
    assert original.import_fingerprint != changed.import_fingerprint


def test_same_bytes_in_different_account_are_not_deduplicated_or_candidate_linked(case):
    content = _xlsx([_spot()])
    first, _ = _import(case, content, _options())
    options = _options(account="synthetic-second-account")
    preview = _preview(case, content, options)
    assert preview.duplicate_batch_id is None
    assert preview.can_import is True
    assert _native(preview)["rows"][0]["status"] == "ready"
    second, events = _import(case, content, options)
    assert first.batch_id != second.batch_id
    assert len(events) == 1
    assert _positions(case) == {(LEDGER_ACCOUNT, "DEMOCOIN"): D("2.5"), ("mexc:" + UID + ":synthetic-second-account", "DEMOCOIN"): D("2.5")}


def test_uid_and_account_boundaries_cannot_collide_in_ledger_namespace(case):
    first_account = "synthetic:spot"
    second_uid, second_account = UID + ":synthetic", "spot"
    _, first_events = _import(case, _xlsx([_spot()]), _options(account=first_account))
    _, second_events = _import(case, _xlsx([_spot(UID=second_uid)]), _options(account=second_account))
    assert len(first_events) == len(second_events) == 1
    assert first_events[0].account != second_events[0].account, "Unterschiedliche UID/Konto-Tupel duerfen nicht durch ':' kollidieren"
    positions = BalanceLedgerBuilder(case[0]).positions()
    assert len(positions) == 2
    assert [p.quantity for p in positions] == [D("2.5"), D("2.5")]


@pytest.mark.parametrize("mismatch", ["missing", "account", "uid", "time", "asset", "quantity", "label", "foreign_source"])
def test_invalid_link_target_is_rejected_before_batch_creation(case, monkeypatch, mismatch):
    changes = {
        "uid": {"UID": "SYNTHETIC-OTHER-UID"},
        "time": {"Creation Time(UTC+01:00)": "2031-07-09 12:00:01"},
        "asset": {"Crypto": "OTHERCOIN"},
        "quantity": {"Quantity": "3.5"},
        "label": {"Transaction Type": "Unknown/Airdrop"},
    }
    options = _options(account="synthetic-other-account") if mismatch == "account" else _options()
    batch, _ = _import(case, _xlsx([_spot(**changes.get(mismatch, {}))]), options)
    root, = _raw(case, batch)
    target = root.raw_object_id
    if mismatch == "missing":
        target = "synthetic-nonexistent-raw-id"
    if mismatch == "foreign_source":
        foreign = case[0].create_import_batch(ImportBatchCreate(source="manual-csv", filename="synthetic.csv", file_hash="synthetic-foreign"))
        case[0].add_raw_objects(foreign.batch_id, [RawObjectCreate(row_number=2, payload_json=root.payload_json)])
        target = case[0].list_raw_objects(foreign.batch_id)[0].raw_object_id
    incoming = _options("funding_other", decisions={"2": _decision("link_existing", target)})
    content = _csv([_funding()])
    preview = _preview(case, content, incoming, filename="synthetic.csv", sheet=None)
    assert preview.can_import is False
    assert _native(preview)["rows"][0]["decision_stale"] is False
    _reject_before_batch(case, monkeypatch, content, incoming, filename="synthetic.csv", sheet=None)


def test_link_to_link_is_rejected_instead_of_chaining_or_cycling(case, monkeypatch):
    _, root = _seed(case)
    first_options = _options("funding_other", decisions={"2": _decision("link_existing", root.raw_object_id)})
    linked_batch, _ = _import(case, _xlsx([_funding()], "funding_other"), first_options)
    linked_raw, = _raw(case, linked_batch)
    options = _options(decisions={"2": _decision("link_existing", linked_raw.raw_object_id)})
    content = _xlsx([_spot()], revision="third")
    preview = _preview(case, content, options)
    assert preview.can_import is False
    assert _native(preview)["rows"][0]["decision_stale"] is False
    _reject_before_batch(case, monkeypatch, content, options)


@pytest.mark.parametrize("kind", ["asset_snapshot", "futures_capital_flow"])
def test_snapshot_or_bonus_cannot_be_link_root(case, monkeypatch, kind):
    rows = [_snapshot_row()] if kind == "asset_snapshot" else _bonus_rows()
    batch, events = _import(case, _xlsx(rows, kind), _options(kind))
    assert events == []
    target = _raw(case, batch)[0].raw_object_id
    options = _options(decisions={"2": _decision("link_existing", target)})
    content = _xlsx([_spot()])
    preview = _preview(case, content, options)
    assert preview.can_import is False
    assert _native(preview)["rows"][0]["decision_stale"] is False
    _reject_before_batch(case, monkeypatch, content, options)


def test_invalid_second_row_decision_leaves_no_partial_first_row_or_batch(case, monkeypatch):
    _, root = _seed(case)
    options = _options("funding_other", decisions={
        "2": _decision("link_existing", root.raw_object_id),
        "3": _decision("link_existing", "synthetic-missing-target"),
    })
    _reject_before_batch(case, monkeypatch, _csv([_funding(), _funding()]), options, filename="synthetic.csv", sheet=None)


@pytest.mark.parametrize("action", ["separate", "link_existing"])
def test_short_decision_reason_blocks_direct_import_before_batch(case, monkeypatch, action):
    _, root = _seed(case)
    options = _options("funding_other", decisions={"2": _decision(action, root.raw_object_id if action == "link_existing" else None, reason="123456789")})
    _reject_before_batch(case, monkeypatch, _xlsx([_funding()], "funding_other"), options)


def test_snapshot_is_only_raw_evidence_and_never_acquisition_or_valuation_cash(case):
    content = _xlsx([_snapshot_row()], "asset_snapshot")
    options = _options("asset_snapshot")
    preview = _preview(case, content, options)
    assert preview.can_import is True
    _counts(_native(preview), raw_rows=1, bookable_rows=0, linked_rows=0, snapshot_rows=1, bonus_rows=0)
    batch, events = _import(case, content, options)
    assert events == []
    payload, meta = _metadata(_raw(case, batch)[0])
    assert meta["role"] == "snapshot"
    assert payload["Date(UTC+8)"] == "20310709"
    assert (payload["Total"], payload["USDT Valuation"]) == ("2.5", "123")
    assert _positions(case) == {}
    assert _state(case)["tax_lots"] == []
    assert _state(case)["tax_lines"] == []
    assert _state(case)["derivative_lines"] == []


def test_bonus_allocation_and_negative_reversal_remain_two_raws_no_events(case):
    rows = _bonus_rows()
    content = _xlsx(rows, "futures_capital_flow")
    options = _options("futures_capital_flow")
    preview = _preview(case, content, options)
    assert preview.can_import is True
    _counts(_native(preview), raw_rows=2, bookable_rows=0, linked_rows=0, snapshot_rows=0, bonus_rows=2)
    batch, events = _import(case, content, options)
    assert events == []
    raws = _raw(case, batch)
    assert [raw.row_number for raw in raws] == [2, 3]
    amounts = []
    for raw, original in zip(raws, rows):
        payload, meta = _metadata(raw)
        assert {key: payload[key] for key in original} == original
        assert meta["role"] == "bonus"
        amounts.append(D(payload["Amount"]))
    assert amounts == [D("7"), D("-7")]
    assert sum(amounts) == D("0")
    assert _positions(case) == {}
    assert _state(case)["tax_lots"] == []
    assert _state(case)["derivative_lines"] == []


@pytest.mark.parametrize("kind", ["futures_statement", "earn_fixed", "earn_flexible", "send_history"])
def test_same_headers_other_products_are_review_evidence_not_spot_bookings(case, kind):
    content = _xlsx([_spot()], kind)
    options = _options(kind)
    preview = _preview(case, content, options)
    assert preview.can_import is True
    native = _native(preview)
    _counts(native, raw_rows=1, bookable_rows=0, review_rows=1)
    assert native["rows"][0]["status"] in {"evidence_only", "review"}
    assert native["rows"][0]["reason"]
    batch, events = _import(case, content, options)
    assert events == []
    assert _metadata(_raw(case, batch)[0])[1]["role"] == "review"
    assert _positions(case) == {}


def test_unknown_other_label_is_raw_review_not_generic_unknown_balance_event(case):
    row = _spot(**{"Transaction Type": "SYNTHETIC_UNDOCUMENTED_OPERATION"})
    content = _xlsx([row])
    preview = _preview(case, content, _options())
    native = _native(preview)
    _counts(native, raw_rows=1, bookable_rows=0, review_rows=1)
    assert native["rows"][0]["status"] == "review"
    assert native["rows"][0]["reason"]
    batch, events = _import(case, content, _options())
    assert events == []
    assert _metadata(_raw(case, batch)[0])[1]["role"] == "review"
    assert _positions(case) == {}


def test_airdrop_has_own_mexc_classification_export_blocker(case):
    from steuerreport_v2.domain.review import ReviewService

    _seed(case)
    gate = ReviewService(case[0]).gate()
    assert gate.allow_export is False
    assert any(issue.blocks_export and issue.code == "mexc_classification_required" for issue in gate.issues), [i.code for i in gate.issues]
    assert [(e.event_type, e.quantity) for e in case[0].list_canonical_events()] == [("unknown", D("2.5"))]


def _string_leaves(value):
    if isinstance(value, dict):
        for item in value.values():
            yield from _string_leaves(item)
    elif isinstance(value, list):
        for item in value:
            yield from _string_leaves(item)
    elif isinstance(value, str):
        yield value


def test_assistant_preview_passes_native_options_and_trace_exposes_both_evidence_directions(case):
    from fastapi.testclient import TestClient
    from steuerreport_v2.api.app import create_app

    content = _xlsx([_spot()])
    request = {"import_mode": "file", **_request(content, _options())}
    before = _state(case)
    with TestClient(create_app(case[0])) as client:
        response = client.post("/api/v1/import/assistant/preview", json=request)
        assert response.status_code == 200
        data = response.json()["data"]
        assert data["can_execute"] is True
        assert data["preview"]["mexc"]["raw_rows"] == 1
        assert data["preview"]["mexc"]["bookable_rows"] == 1
        assert _state(case) == before
        root_response = client.post("/api/v1/import/assistant/execute", json=request)
        assert root_response.status_code == 200
        root_result = root_response.json()
        assert root_result["status"] == "ok", root_result
        root_batch_id = root_result["data"]["batch"]["batch_id"]
        root, = case[0].list_raw_objects(root_batch_id, limit=None)
        funding = _xlsx([_funding()], "funding_other")
        linked_options = _options("funding_other", decisions={"2": _decision("link_existing", root.raw_object_id)})
        _bind_decisions(case, funding, linked_options)
        linked_request = {"import_mode": "file", **_request(funding, linked_options)}
        linked_response = client.post("/api/v1/import/assistant/execute", json=linked_request)
        assert linked_response.status_code == 200
        linked_result = linked_response.json()
        assert linked_result["status"] == "ok", linked_result
        linked_batch_id = linked_result["data"]["batch"]["batch_id"]
        linked_raw, = case[0].list_raw_objects(linked_batch_id, limit=None)
        for batch_id in (root_batch_id, linked_batch_id):
            trace_response = client.get(f"/api/v1/import/batches/{batch_id}/trace.json")
            assert trace_response.status_code == 200
            trace = trace_response.json()
            assert trace["batch"]["batch_id"] == batch_id
            assert trace["raw_objects"]
            assert "canonical_events" in trace
            assert trace["mexc_evidence"], "Belegtrace braucht Vorwaerts- und Rueckverweise"
            ids = set(_string_leaves(trace["mexc_evidence"]))
            assert {root.raw_object_id, linked_raw.raw_object_id} <= ids
        assert len(case[0].list_canonical_events()) == 1
        assert _positions(case) == {(LEDGER_ACCOUNT, "DEMOCOIN"): D("2.5")}


@pytest.mark.parametrize("id_column", ["Transaction ID", "Record ID"])
@pytest.mark.parametrize("layout", ["same_file", "separate_files", "cross_export"])
def test_distinct_provider_ids_preserve_two_legitimate_equal_bookings(case, id_column, layout):
    first_id, second_id = "SYNTHETIC-TXN-A", "SYNTHETIC-TXN-B"
    first_row = _spot(**{id_column: first_id})
    first_content = _xlsx([first_row], extra_headers=(id_column,))
    if layout == "same_file":
        content = _xlsx([first_row, _spot(**{id_column: second_id})], extra_headers=(id_column,))
        preview = _preview(case, content, _options())
        batches = [_import(case, content, _options())[0]]
    else:
        first_batch, _ = _import(case, first_content, _options())
        kind = "funding_other" if layout == "cross_export" else "spot_statement"
        second_row = _funding(**{id_column: second_id}) if layout == "cross_export" else _spot(**{id_column: second_id})
        content = _xlsx([second_row], kind, extra_headers=(id_column,))
        preview = _preview(case, content, _options(kind))
        batches = [first_batch, _import(case, content, _options(kind))[0]]
    assert preview.can_import is True
    native = _native(preview)
    _counts(native, raw_rows=2 if layout == "same_file" else 1,
            bookable_rows=2 if layout == "same_file" else 1, linked_rows=0)
    assert all(row["status"] == "ready" and row["candidates"] == [] for row in native["rows"])
    raws = [raw for batch in batches for raw in _raw(case, batch)]
    assert len(raws) == 2
    assert {_metadata(raw)[1]["provider_id"] for raw in raws} == {first_id, second_id}
    assert {_metadata(raw)[0][id_column] for raw in raws} == {first_id, second_id}
    events = case[0].list_canonical_events()
    assert len(events) == 2
    assert all(event.quantity == D("2.5") and event.event_type == "unknown" for event in events)
    assert len({event.raw_object_id for event in events}) == 2
    assert _positions(case) == {(LEDGER_ACCOUNT, "DEMOCOIN"): D("5.0")}


@pytest.mark.parametrize("id_column", ["Transaction ID", "Record ID"])
@pytest.mark.parametrize("mismatch", ["quantity", "time", "asset", "direction"])
@pytest.mark.parametrize("action", [None, "separate", "link_existing"])
def test_same_provider_id_conflict_blocks_even_with_decision_before_batch(case, monkeypatch, id_column, mismatch, action):
    provider_id = "SYNTHETIC-CONFLICT-ID"
    first = _xlsx([_spot(**{id_column: provider_id})], extra_headers=(id_column,))
    root_batch, _ = _import(case, first, _options())
    root, = _raw(case, root_batch)
    changes = {
        "quantity": {"Quantity": "3.5"},
        "time": {"Creation Time(UTC+01:00)": "2031-07-09 12:00:01"},
        "asset": {"Crypto": "OTHERCOIN"},
        "direction": {"Direction": "Outflow"},
    }[mismatch]
    content = _xlsx([_spot(**{id_column: provider_id, **changes})], extra_headers=(id_column,))
    options = _options()
    if action:
        options["decisions"] = {"2": _decision(action, root.raw_object_id if action == "link_existing" else None)}
    before = _state(case)
    preview = _preview(case, content, options)
    assert preview.can_import is False
    native = _native(preview)
    _counts(native, raw_rows=1, bookable_rows=0, linked_rows=0)
    assert native["blocking_reasons"]
    assert native["rows"][0]["status"] == "invalid"
    assert native["rows"][0]["reason"]
    assert _state(case) == before
    _reject_before_batch(case, monkeypatch, content, options)
    assert _positions(case) == {(LEDGER_ACCOUNT, "DEMOCOIN"): D("2.5")}


@pytest.mark.parametrize("id_column", ["Transaction ID", "Record ID"])
def test_provider_id_conflict_in_second_physical_row_cannot_partially_import_first(case, monkeypatch, id_column):
    content = _xlsx([
        _spot(**{id_column: "SYNTHETIC-ONE-ID"}),
        _spot(**{id_column: "SYNTHETIC-ONE-ID", "Quantity": "3.5"}),
    ], extra_headers=(id_column,))
    preview = _preview(case, content, _options())
    assert preview.can_import is False
    assert _native(preview)["rows"][1]["row_number"] == 3
    assert _native(preview)["rows"][1]["status"] == "invalid"
    _reject_before_batch(case, monkeypatch, content, _options())
    assert case[0].list_import_batches() == []
    assert case[0].list_canonical_events() == []


@pytest.mark.parametrize("id_column", ["Transaction ID", "Record ID"])
def test_same_provider_id_across_header_aliases_links_one_event_two_evidence_rows(case, monkeypatch, id_column):
    provider_id = "SYNTHETIC-SHARED-ID"
    first = _xlsx([_spot(**{id_column: provider_id})], extra_headers=(id_column,))
    root_batch, _ = _import(case, first, _options())
    root, = _raw(case, root_batch)
    other_column = "Record ID" if id_column == "Transaction ID" else "Transaction ID"
    content = _xlsx([_funding(**{other_column: provider_id})], "funding_other", extra_headers=(other_column,))
    options = _options("funding_other")
    preview = _preview(case, content, options)
    assert preview.can_import is False
    row, = _native(preview)["rows"]
    assert row["status"] == "duplicate_candidate"
    assert root.raw_object_id in set(_string_leaves(row["candidates"]))
    _reject_before_batch(case, monkeypatch, content, options)
    options["decisions"] = {"2": _decision("link_existing", root.raw_object_id)}
    batch, events = _import(case, content, options)
    assert events == []
    linked, = _raw(case, batch)
    payload, meta = _metadata(linked)
    assert payload[other_column] == provider_id
    assert meta["provider_id"] == provider_id
    assert meta["link_target_raw_object_id"] == root.raw_object_id
    assert meta["role"] == "linked"
    assert len(case[0].list_canonical_events()) == 1
    assert _positions(case) == {(LEDGER_ACCOUNT, "DEMOCOIN"): D("2.5")}


@pytest.mark.parametrize("id_column", ["Transaction ID", "Record ID"])
@pytest.mark.parametrize("boundary", ["account", "uid"])
def test_provider_id_scope_is_uid_and_account_not_global(case, id_column, boundary):
    provider_id = "SYNTHETIC-ID-REUSED-IN-OTHER-CONTEXT"
    content = _xlsx([_spot(**{id_column: provider_id})], extra_headers=(id_column,))
    first_batch, _ = _import(case, content, _options())
    second_options = _options(account="synthetic-other-account") if boundary == "account" else _options()
    other_uid = "SYNTHETIC-OTHER-UID" if boundary == "uid" else UID
    second_content = _xlsx([_spot(**{id_column: provider_id, "UID": other_uid, "Quantity": "3.5"})], extra_headers=(id_column,))
    preview = _preview(case, second_content, second_options)
    assert preview.can_import is True
    assert _native(preview)["rows"][0]["status"] == "ready"
    assert _native(preview)["rows"][0]["candidates"] == []
    second_batch, _ = _import(case, second_content, second_options)
    assert first_batch.batch_id != second_batch.batch_id
    events = case[0].list_canonical_events()
    assert len(events) == 2
    assert len({event.account for event in events}) == 2
    assert sorted(_positions(case).values()) == [D("2.5"), D("3.5")]


@pytest.mark.parametrize("generic_id_column", [False, True])
@pytest.mark.parametrize("second_quantity,expected_balance", [("2.5", "5.0"), ("3.5", "6.0")])
def test_uid_or_generic_id_equal_uid_is_never_a_provider_transaction_id(case, generic_id_column, second_quantity, expected_balance):
    extra = ("ID",) if generic_id_column else ()
    arbitrary_id = {"ID": UID} if generic_id_column else {}
    first = _xlsx([_spot(**arbitrary_id)], extra_headers=extra)
    first_batch, _ = _import(case, first, _options())
    second = _xlsx([_spot(**{**arbitrary_id, "Quantity": second_quantity})], extra_headers=extra, revision="second-export")
    options = _options()
    preview = _preview(case, second, options)
    row, = _native(preview)["rows"]
    if second_quantity == "2.5":
        assert row["status"] == "duplicate_candidate"
        options["decisions"] = {"2": _decision("separate")}
    else:
        assert preview.can_import is True
        assert row["status"] == "ready"
        assert row["candidates"] == []
    second_batch, _ = _import(case, second, options)
    for batch in (first_batch, second_batch):
        raw, = _raw(case, batch)
        payload, meta = _metadata(raw)
        assert payload["UID"] == UID
        assert meta["uid"] == UID
        assert meta.get("provider_id") in (None, ""), "UID ist Kontokennung, keine Provider-Buchungs-ID"
        if generic_id_column:
            assert payload["ID"] == UID
    assert len(case[0].list_canonical_events()) == 2
    assert _positions(case) == {(LEDGER_ACCOUNT, "DEMOCOIN"): D(expected_balance)}


@pytest.mark.parametrize("legacy_meta", [None, "{}", "not-json", '{"version":0,"role":"book"}'])
@pytest.mark.parametrize("decision", [False, True])
def test_legacy_mexc_without_native_binding_blocks_before_batch(case, monkeypatch, legacy_meta, decision):
    legacy = case[0].create_import_batch(ImportBatchCreate(source="mexc", filename="synthetic-legacy.xlsx", file_hash="synthetic-legacy-hash"))
    payload = _spot()
    if legacy_meta is not None:
        payload["_mexc"] = legacy_meta
    case[0].add_raw_objects(legacy.batch_id, [RawObjectCreate(row_number=2, payload_json=json.dumps(payload, sort_keys=True))])
    original, = _raw(case, legacy)
    before = _state(case)
    content = _xlsx([_spot()])
    options = _options(decisions={"2": _decision("separate")}) if decision else _options()
    try:
        preview = _preview(case, content, options)
    except ValueError as error:
        assert str(error)
    else:
        assert preview.can_import is False
        assert _native(preview)["blocking_reasons"]
    assert _state(case) == before
    _reject_before_batch(case, monkeypatch, content, options)
    assert _raw(case, legacy)[0].payload_json == original.payload_json
    assert len(case[0].list_import_batches()) == 1
    assert case[0].list_canonical_events() == []


def test_foreign_source_legacy_raw_does_not_block_native_mexc(case):
    legacy = case[0].create_import_batch(ImportBatchCreate(source="manual-csv", filename="synthetic-foreign.csv", file_hash="synthetic-foreign-legacy"))
    case[0].add_raw_objects(legacy.batch_id, [RawObjectCreate(row_number=2, payload_json=json.dumps(_spot()))])
    content = _xlsx([_spot()])
    preview = _preview(case, content, _options())
    assert preview.can_import is True
    assert _native(preview)["rows"][0]["candidates"] == []
    batch, events = _import(case, content, _options())
    assert batch.batch_id != legacy.batch_id
    assert len(events) == 1
    assert _positions(case) == {(LEDGER_ACCOUNT, "DEMOCOIN"): D("2.5")}


def test_37_digit_running_balance_keeps_single_atto_unit_changes_and_replay(case):
    rows = [
        _spot(Quantity="1234567890123456789.123456789012345678"),
        _spot(**{"Creation Time(UTC+01:00)": "2031-07-09 12:00:01", "Quantity": "0.000000000000000001"}),
        _spot(**{"Creation Time(UTC+01:00)": "2031-07-09 12:00:02", "Direction": "Outflow", "Quantity": "0.000000000000000002"}),
    ]
    batch, events = _import(case, _xlsx(rows), _options())
    assert len(events) == 3
    raw_before = [(raw.raw_object_id, raw.payload_json) for raw in _raw(case, batch)]
    expected = [
        (D("1234567890123456789.123456789012345678"), D("1234567890123456789.123456789012345678")),
        (D("0.000000000000000001"), D("1234567890123456789.123456789012345679")),
        (D("-0.000000000000000002"), D("1234567890123456789.123456789012345677")),
    ]
    for replay in (False, True):
        if replay:
            assert len(RawNormalizer(case[0]).normalize_batch(batch.batch_id)) == 3
            BalanceLedgerBuilder(case[0]).rebuild()
        with case[0].connect() as con:
            actual = con.execute("select delta_quantity,balance_quantity from balance_ledger order by event_time").fetchall()
        assert actual == expected
        assert _positions(case) == {(LEDGER_ACCOUNT, "DEMOCOIN"): D("1234567890123456789.123456789012345677")}
        assert [(raw.raw_object_id, raw.payload_json) for raw in _raw(case, batch)] == raw_before


def test_binding_helper_reads_preview_without_decisions_and_keeps_existing_token(case, monkeypatch):
    _seed(case)
    content = _xlsx([_funding(), _funding()], "funding_other")
    supplied = "0" * 64
    options = _options("funding_other", decisions={
        "2": {**_decision("separate"), "input_digest": supplied},
        "3": _decision("separate"),
    })
    original_preview = case[1].preview_base64
    calls = []

    def observed_preview(**request):
        calls.append(request["mexc"])
        assert not request["mexc"].get("decisions")
        return original_preview(**request)

    before = _state(case)
    monkeypatch.setattr(case[1], "preview_base64", observed_preview)
    assert _bind_decisions(case, content, options) is options
    assert len(calls) == 1
    assert options["decisions"]["2"]["input_digest"] == supplied
    assert re.fullmatch(r"[0-9a-fA-F]{64}", options["decisions"]["3"]["input_digest"])
    assert _state(case) == before


@pytest.mark.parametrize("supplied", [None, "", "0" * 64])
def test_binding_helper_never_replaces_explicit_token_even_when_invalid(case, monkeypatch, supplied):
    options = _options(decisions={"2": {**_decision("separate"), "input_digest": supplied}})

    def forbidden(**request):
        pytest.fail("Ein expliziter Token darf nicht durch erneute Vorschau ersetzt werden")

    monkeypatch.setattr(case[1], "preview_base64", forbidden)
    _bind_decisions(case, b"not read for already supplied token", options)
    assert options["decisions"]["2"]["input_digest"] == supplied


@pytest.mark.parametrize("action", ["link_existing", "separate"])
@pytest.mark.parametrize("explicit_null", [False, True])
def test_missing_decision_digest_blocks_direct_import_before_batch(case, monkeypatch, action, explicit_null):
    _, root = _seed(case)
    content = _xlsx([_funding()], "funding_other")
    options = _options("funding_other", decisions={"2": _decision(action, root.raw_object_id if action == "link_existing" else None)})
    if explicit_null:
        options["decisions"]["2"]["input_digest"] = None
    preview = _preview(case, content, options, bind_decisions=False)
    assert preview.can_import is False
    assert _native(preview)["rows"][0]["reason"]
    _reject_before_batch(case, monkeypatch, content, options, bind_decisions=False)
    if explicit_null:
        assert options["decisions"]["2"]["input_digest"] is None
    else:
        assert "input_digest" not in options["decisions"]["2"]


@pytest.mark.parametrize("action", ["link_existing", "separate"])
def test_explicit_wrong_decision_digest_is_not_silently_rebound(case, monkeypatch, action):
    _, root = _seed(case)
    content = _xlsx([_funding()], "funding_other")
    supplied = "0" * 64
    options = _options("funding_other", decisions={"2": {
        **_decision(action, root.raw_object_id if action == "link_existing" else None),
        "input_digest": supplied,
    }})
    preview = _preview(case, content, options)
    assert preview.can_import is False
    assert _native(preview)["rows"][0]["decision_stale"] is True
    _reject_before_batch(case, monkeypatch, content, options)
    assert options["decisions"]["2"]["input_digest"] == supplied


@pytest.mark.parametrize("action", ["link_existing", "separate"])
@pytest.mark.parametrize("scope", ["bytes", "sheet", "account", "timezone"])
def test_bound_decision_rejects_changed_original_bytes_or_import_scope(case, monkeypatch, action, scope):
    _, root = _seed(case)
    content = _xlsx([_funding()], "funding_other", sheets=("Sheet1", "Sheet2"))
    options = _options("funding_other", decisions={"2": _decision(action, root.raw_object_id if action == "link_existing" else None)})
    _bind_decisions(case, content, options)
    claimed = options["decisions"]["2"]["input_digest"]
    changed = content
    kwargs = {}
    if scope == "bytes":
        changed = _xlsx([_funding()], "funding_other", sheets=("Sheet1", "Sheet2"), revision="changed-original-bytes")
        assert hashlib.sha256(changed).digest() != hashlib.sha256(content).digest()
    elif scope == "sheet":
        kwargs["sheet"] = "Sheet2"
    elif scope == "account":
        options["account"] = "synthetic-other-account"
    else:
        options["timezone"] = "UTC"
    unclaimed = {key: value for key, value in options.items() if key != "decisions"}
    fresh = _preview(case, changed, unclaimed, **kwargs)
    assert _native(fresh)["rows"][0]["input_digest"] != claimed
    preview = _preview(case, changed, options, **kwargs)
    assert preview.can_import is False
    stale = _native(preview)["rows"][0]["decision_stale"]
    _reject_before_batch(case, monkeypatch, changed, options, **kwargs)
    assert options["decisions"]["2"]["input_digest"] == claimed
    assert _positions(case) == {(LEDGER_ACCOUNT, "DEMOCOIN"): D("2.5")}
    assert stale is True


@pytest.mark.parametrize("action", ["link_existing", "separate"])
def test_bound_decision_rejects_changed_candidate_roots(case, monkeypatch, action):
    _, root = _seed(case)
    content = _xlsx([_funding()], "funding_other")
    options = _options("funding_other", decisions={"2": _decision(action, root.raw_object_id if action == "link_existing" else None)})
    _bind_decisions(case, content, options)
    claimed = options["decisions"]["2"]["input_digest"]
    additional = _xlsx([_spot()], revision="another-legitimate-booking")
    _import(case, additional, _options(decisions={"2": _decision("separate")}))
    fresh = _preview(case, content, _options("funding_other"))
    assert len(_native(fresh)["rows"][0]["candidates"]) == 2
    assert _native(fresh)["rows"][0]["input_digest"] != claimed
    preview = _preview(case, content, options)
    assert preview.can_import is False
    assert _native(preview)["rows"][0]["decision_stale"] is True
    _reject_before_batch(case, monkeypatch, content, options)
    assert options["decisions"]["2"]["input_digest"] == claimed
    assert _positions(case) == {(LEDGER_ACCOUNT, "DEMOCOIN"): D("5.0")}


def test_decision_digest_cannot_be_moved_between_equal_physical_rows(case, monkeypatch):
    _seed(case)
    content = _xlsx([_funding(), _funding()], "funding_other")
    unclaimed = _preview(case, content, _options("funding_other"))
    rows = _native(unclaimed)["rows"]
    assert [row["row_number"] for row in rows] == [2, 3]
    assert rows[0]["input_digest"] != rows[1]["input_digest"]
    options = _options("funding_other", decisions={
        "2": _decision("separate"),
        "3": {**_decision("separate"), "input_digest": rows[0]["input_digest"]},
    })
    preview = _preview(case, content, options)
    assert preview.can_import is False
    assert _native(preview)["rows"][1]["decision_stale"] is True
    _reject_before_batch(case, monkeypatch, content, options)
    assert options["decisions"]["3"]["input_digest"] == rows[0]["input_digest"]
    assert _positions(case) == {(LEDGER_ACCOUNT, "DEMOCOIN"): D("2.5")}


@pytest.mark.parametrize("id_column", ["Transaction ID", "Record ID"])
def test_same_provider_id_cannot_be_separate_even_with_fresh_digest(case, monkeypatch, id_column):
    provider_id = "SYNTHETIC-UNIQUE-PROVIDER-ID"
    first = _xlsx([_spot(**{id_column: provider_id})], extra_headers=(id_column,))
    root_batch, _ = _import(case, first, _options())
    root, = _raw(case, root_batch)
    content = _xlsx([_funding(**{id_column: provider_id})], "funding_other", extra_headers=(id_column,))
    options = _options("funding_other", decisions={"2": _decision("separate")})
    preview = _preview(case, content, options)
    assert preview.can_import is False
    native = _native(preview)
    _counts(native, raw_rows=1, bookable_rows=0, linked_rows=0)
    assert native["rows"][0]["decision_stale"] is False
    _reject_before_batch(case, monkeypatch, content, options)
    options["decisions"]["2"].update(action="link_existing", target_raw_object_id=root.raw_object_id)
    batch, events = _import(case, content, options)
    assert events == []
    assert _metadata(_raw(case, batch)[0])[1]["role"] == "linked"
    assert len(case[0].list_canonical_events()) == 1
    assert _positions(case) == {(LEDGER_ACCOUNT, "DEMOCOIN"): D("2.5")}


@pytest.mark.parametrize("id_column", ["Transaction ID", "Record ID"])
@pytest.mark.parametrize("action", [None, "separate", "link_existing"])
def test_provider_id_introduced_by_linked_evidence_also_blocks_conflicts(case, monkeypatch, id_column, action):
    _, root = _seed(case)
    provider_id = "SYNTHETIC-ID-FROM-LINKED-EVIDENCE"
    linked_content = _xlsx([_funding(**{id_column: provider_id})], "funding_other", extra_headers=(id_column,))
    linked_options = _options("funding_other", decisions={"2": _decision("link_existing", root.raw_object_id)})
    linked_batch, _ = _import(case, linked_content, linked_options)
    assert _metadata(_raw(case, linked_batch)[0])[1]["provider_id"] == provider_id
    conflict = _xlsx([_spot(**{id_column: provider_id, "Quantity": "3.5"})], extra_headers=(id_column,))
    options = _options()
    if action:
        options["decisions"] = {"2": _decision(action, root.raw_object_id if action == "link_existing" else None)}
    preview = _preview(case, conflict, options)
    assert preview.can_import is False
    assert _native(preview)["rows"][0]["decision_stale"] is False
    _reject_before_batch(case, monkeypatch, conflict, options)
    assert len(case[0].list_canonical_events()) == 1
    assert _positions(case) == {(LEDGER_ACCOUNT, "DEMOCOIN"): D("2.5")}


@pytest.mark.parametrize("action", ["link_existing", "separate"])
def test_linked_provider_evidence_changes_candidate_digest_even_when_root_id_stays(case, monkeypatch, action):
    _, root = _seed(case)
    provider_id = "SYNTHETIC-LATER-LINKED-PROVIDER-ID"
    content = _xlsx([_funding(**{"Transaction ID": provider_id})], "funding_other", extra_headers=("Transaction ID",), revision="pending-decision")
    options = _options("funding_other", decisions={"2": _decision(action, root.raw_object_id if action == "link_existing" else None)})
    _bind_decisions(case, content, options)
    claimed = options["decisions"]["2"]["input_digest"]
    new_evidence = _xlsx([_funding(**{"Transaction ID": provider_id})], "funding_other", extra_headers=("Transaction ID",), revision="new-linked-evidence")
    _import(case, new_evidence, _options("funding_other", decisions={"2": _decision("link_existing", root.raw_object_id)}))
    unclaimed = _preview(case, content, _options("funding_other"))
    row, = _native(unclaimed)["rows"]
    assert root.raw_object_id in set(_string_leaves(row["candidates"]))
    assert row["input_digest"] != claimed
    preview = _preview(case, content, options)
    assert preview.can_import is False
    assert _native(preview)["rows"][0]["decision_stale"] is True
    _reject_before_batch(case, monkeypatch, content, options)
    assert _positions(case) == {(LEDGER_ACCOUNT, "DEMOCOIN"): D("2.5")}


def test_typed_decision_model_gets_digest_only_when_field_was_not_explicit(case, monkeypatch):
    from steuerreport_v2.contracts.mexc_import import MexcImportOptions

    _seed(case)
    content = _xlsx([_funding()], "funding_other")
    missing = MexcImportOptions.model_validate(_options("funding_other", decisions={"2": _decision("separate")}))
    assert "input_digest" not in missing.decisions["2"].model_fields_set
    _bind_decisions(case, content, missing)
    assert re.fullmatch(r"[0-9a-f]{64}", missing.decisions["2"].input_digest)
    explicit = MexcImportOptions.model_validate(_options("funding_other", decisions={"2": {**_decision("separate"), "input_digest": None}}))
    assert "input_digest" in explicit.decisions["2"].model_fields_set
    _bind_decisions(case, content, explicit)
    assert explicit.decisions["2"].input_digest is None
    _reject_before_batch(case, monkeypatch, content, explicit)
    _, events = _import(case, content, missing)
    assert len(events) == 1
    assert _positions(case) == {(LEDGER_ACCOUNT, "DEMOCOIN"): D("5.0")}


def test_snapshot_only_has_specific_reconciliation_gate_without_fake_history(case):
    from steuerreport_v2.domain.review import ReviewService

    content = _xlsx([_snapshot_row()], "asset_snapshot")
    batch, events = _import(case, content, _options("asset_snapshot"))
    assert events == []
    assert batch.batch_id not in case[0].unmapped_raw_object_counts()
    before = _state(case)
    gate = ReviewService(case[0]).gate()
    assert gate.allow_export is False
    assert any(i.code == "mexc_snapshot_reconciliation_required" and i.blocks_export for i in gate.issues)
    assert not any(i.code == "mexc_original_integrity_failed" for i in gate.issues)
    assert _state(case) == before
    assert _positions(case) == {}


@pytest.mark.parametrize("damage", ["missing", "hash", "same_size_mtime", "symlink"])
def test_original_integrity_gate_detects_damage_and_recovers_only_after_exact_restore(case, damage):
    from steuerreport_v2.domain.review import ReviewService

    content = _xlsx([_snapshot_row()], "asset_snapshot")
    batch, _ = _import(case, content, _options("asset_snapshot"))
    raw, = _raw(case, batch)
    original_payload = raw.payload_json
    archive = case[0].path.parent / "mexc_originals" / (hashlib.sha256(content).hexdigest() + ".original")
    assert archive.read_bytes() == content
    baseline = ReviewService(case[0]).gate()
    assert not any(i.code == "mexc_original_integrity_failed" for i in baseline.issues)
    stamp = archive.stat()
    if damage == "missing":
        archive.unlink()
    elif damage == "hash":
        archive.write_bytes(b"SYNTHETIC CORRUPTED ORIGINAL")
    elif damage == "same_size_mtime":
        changed = bytearray(content)
        changed[-1] ^= 1
        archive.write_bytes(changed)
        os.utime(archive, ns=(stamp.st_atime_ns, stamp.st_mtime_ns))
        assert archive.stat().st_size == stamp.st_size
        assert archive.stat().st_mtime_ns == stamp.st_mtime_ns
    else:
        target = archive.with_suffix(".synthetic-backup")
        archive.rename(target)
        archive.symlink_to(target.name)
        assert archive.read_bytes() == content
    damaged = ReviewService(case[0]).gate()
    assert damaged.allow_export is False
    assert any(i.code == "mexc_original_integrity_failed" and i.blocks_export for i in damaged.issues)
    assert _raw(case, batch)[0].payload_json == original_payload
    if archive.is_symlink():
        archive.unlink()
    archive.write_bytes(content)
    restored = ReviewService(case[0]).gate()
    assert not any(i.code == "mexc_original_integrity_failed" for i in restored.issues)
    assert restored.allow_export is False
    assert any(i.code == "mexc_snapshot_reconciliation_required" and i.blocks_export for i in restored.issues)
    assert _positions(case) == {}
    assert case[0].list_canonical_events() == []


def _replace_metadata(case, raw, meta):
    payload = json.loads(raw.payload_json)
    payload["_mexc"] = json.dumps(meta, sort_keys=True)
    with case[0].connect() as con:
        con.execute("update raw_objects set payload_json=? where raw_object_id=?", [json.dumps(payload, sort_keys=True), raw.raw_object_id])


@pytest.mark.parametrize("forged_role", ["book", "linked", "bonus", "review", "not-a-native-role"])
def test_snapshot_cannot_escape_binding_gate_by_changing_metadata_role(case, forged_role):
    from steuerreport_v2.domain.review import ReviewService

    batch, _ = _import(case, _xlsx([_snapshot_row()], "asset_snapshot"), _options("asset_snapshot"))
    raw, = _raw(case, batch)
    _, meta = _metadata(raw)
    meta["role"] = forged_role
    _replace_metadata(case, raw, meta)
    assert RawNormalizer(case[0]).normalize_batch(batch.batch_id) == []
    BalanceLedgerBuilder(case[0]).rebuild()
    gate = ReviewService(case[0]).gate()
    assert gate.allow_export is False
    assert any(i.code == "mexc_native_binding_invalid" and i.blocks_export for i in gate.issues)
    assert _positions(case) == {}


@pytest.mark.parametrize("action", ["link_existing", "separate"])
def test_persisted_decision_digest_and_evidence_are_verified_not_tax_release(case, action):
    from steuerreport_v2.domain.review import ReviewService

    _, root = _seed(case)
    content = _xlsx([_funding()], "funding_other")
    options = _options("funding_other", decisions={"2": _decision(action, root.raw_object_id if action == "link_existing" else None)})
    batch, events = _import(case, content, options)
    assert len(events) == (0 if action == "link_existing" else 1)
    raw, = _raw(case, batch)
    _, meta = _metadata(raw)
    assert meta["decision_input_digest"] == options["decisions"]["2"]["input_digest"]
    assert root.raw_object_id in set(_string_leaves(meta["decision_evidence"]))
    gate = ReviewService(case[0]).gate()
    assert gate.allow_export is False
    assert any(i.code == "mexc_classification_required" and i.blocks_export for i in gate.issues)
    assert not any(i.code == "mexc_native_binding_invalid" for i in gate.issues)
    meta["decision_input_digest"] = "0" * 64
    _replace_metadata(case, raw, meta)
    assert RawNormalizer(case[0]).normalize_batch(batch.batch_id) == []
    BalanceLedgerBuilder(case[0]).rebuild()
    invalid = ReviewService(case[0]).gate()
    assert invalid.allow_export is False
    assert any(i.code == "mexc_native_binding_invalid" and i.blocks_export for i in invalid.issues)
    assert _positions(case) == {(LEDGER_ACCOUNT, "DEMOCOIN"): D("2.5")}


FUNDING_REVIEW_STATUSES = ("Failed", "Pending", "Completed", "Success", "SYNTHETIC_UNDOCUMENTED_STATUS")


def _funding_status_input(file_format, status):
    row = _funding(Status=status)
    if file_format == "xlsx":
        return _xlsx([row], "funding_other"), {}
    return _csv([row]), {"filename": "synthetic.csv", "sheet": None}


@pytest.mark.parametrize("file_format", ["xlsx", "csv"])
@pytest.mark.parametrize("status", FUNDING_REVIEW_STATUSES)
def test_funding_status_guard_nonempty_is_raw_review_without_events_ledger_or_lots(case, file_format, status):
    from steuerreport_v2.domain.review import ReviewService
    from steuerreport_v2.domain.tax import TaxService

    content, kwargs = _funding_status_input(file_format, status)
    options = _options("funding_other")
    before = _state(case)
    preview = _preview(case, content, options, **kwargs)
    assert preview.can_import is True
    native = _native(preview)
    _counts(native, raw_rows=1, bookable_rows=0, linked_rows=0, review_rows=1, snapshot_rows=0, bonus_rows=0)
    row, = native["rows"]
    assert row["row_number"] == 2
    assert row["status"] == "review"
    assert "status" in row["reason"].lower()
    assert row["candidates"] == []
    assert _state(case) == before
    batch, events = _import(case, content, options, **kwargs)
    assert events == []
    raw, = _raw(case, batch)
    payload, meta = _metadata(raw)
    assert raw.row_number == 2
    assert payload["Status"] == status
    assert payload["Type"] == "Airdrop"
    assert payload["Quantity"] == "2.5"
    assert meta["role"] == "review"
    assert not meta.get("link_target_raw_object_id")
    assert case[0].list_canonical_events() == []
    assert _positions(case) == {}
    TaxService(case[0]).rebuild_fifo()
    state = _state(case)
    assert state["balance_ledger"] == []
    assert state["tax_lots"] == []
    assert state["tax_lines"] == []
    assert case[0].unmapped_raw_object_counts().get(batch.batch_id) == 1
    gate = ReviewService(case[0]).gate()
    assert gate.allow_export is False
    assert any(i.blocks_export and batch.batch_id in i.scope for i in gate.issues)
    original = case[0].path.parent / "mexc_originals" / (hashlib.sha256(content).hexdigest() + ".original")
    assert original.read_bytes() == content


@pytest.mark.parametrize("file_format", ["xlsx", "csv"])
@pytest.mark.parametrize("status", [None, "", "  ", " \t "])
def test_funding_status_guard_empty_or_whitespace_preserves_bookable_unknown_control(case, file_format, status):
    content, kwargs = _funding_status_input(file_format, status)
    options = _options("funding_other")
    preview = _preview(case, content, options, **kwargs)
    assert preview.can_import is True
    native = _native(preview)
    _counts(native, raw_rows=1, bookable_rows=1, linked_rows=0, review_rows=1)
    assert native["rows"][0]["status"] == "ready"
    batch, events = _import(case, content, options, **kwargs)
    assert len(events) == 1
    assert (events[0].event_type, events[0].quantity) == ("unknown", D("2.5"))
    assert events[0].event_time.astimezone(timezone.utc) == WHEN
    assert _positions(case) == {(LEDGER_ACCOUNT, "DEMOCOIN"): D("2.5")}
    raw, = _raw(case, batch)
    payload, meta = _metadata(raw)
    assert meta["role"] == "book"
    assert payload["Status"] == (status if status else None if file_format == "xlsx" else "")


@pytest.mark.parametrize("file_format", ["xlsx", "csv"])
@pytest.mark.parametrize("status", FUNDING_REVIEW_STATUSES)
def test_funding_status_guard_fresh_candidate_link_cannot_override_nonempty_status(case, monkeypatch, file_format, status):
    _, root = _seed(case)
    before = _state(case)
    content, kwargs = _funding_status_input(file_format, status)
    unclaimed = _preview(case, content, _options("funding_other"), **kwargs)
    native = _native(unclaimed)
    _counts(native, raw_rows=1, bookable_rows=0, linked_rows=0, review_rows=1)
    assert native["rows"][0]["status"] == "review"
    assert native["rows"][0]["candidates"] == []
    options = _options("funding_other", decisions={"2": _decision("link_existing", root.raw_object_id)})
    _bind_decisions(case, content, options, **kwargs)
    assert options["decisions"]["2"]["input_digest"] == native["rows"][0]["input_digest"]
    preview = _preview(case, content, options, **kwargs)
    assert preview.can_import is False
    _counts(_native(preview), bookable_rows=0, linked_rows=0)
    assert _native(preview)["rows"][0]["decision_stale"] is False
    _reject_before_batch(case, monkeypatch, content, options, **kwargs)
    assert _state(case) == before
    assert len(case[0].list_canonical_events()) == 1
    assert _positions(case) == {(LEDGER_ACCOUNT, "DEMOCOIN"): D("2.5")}
