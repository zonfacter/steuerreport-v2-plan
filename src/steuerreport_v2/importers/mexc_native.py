"""MEXC-Dateibelege mit expliziter Zuordnung; keine automatische Steuerentscheidung."""
from __future__ import annotations

from collections import defaultdict
import csv
from datetime import datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation, localcontext
from io import BytesIO, StringIO
import hashlib
import json
import os
from pathlib import Path
import posixpath
import re
from threading import RLock
from urllib.parse import quote
from xml.etree import ElementTree as ET
from zoneinfo import ZoneInfo
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

from steuerreport_v2.contracts.events import CanonicalEventCreate
from steuerreport_v2.contracts.imports import CsvImportPreview, ImportBatchCreate, RawObjectCreate
from steuerreport_v2.contracts.mexc_import import MexcCandidate, MexcImportOptions, MexcImportPreview, MexcRowPreview
from steuerreport_v2.db.duckdb_store import DuckDbStore

LOCK = RLock()
MAX_ROWS = 10000
STATEMENT_KINDS = {"spot_statement", "futures_statement", "earn_fixed", "earn_flexible", "send_history"}
NATIVE_KINDS = {"spot_statement", "funding_other"}
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def metadata(row: dict) -> dict:
    try:
        value = json.loads(row.get("_mexc") or "{}")
        return value if isinstance(value, dict) else {}
    except (ValueError, TypeError):
        return {}


def _hash(value) -> str:
    return hashlib.sha256(json.dumps(value, sort_keys=True, separators=(",", ":")).encode()).hexdigest()


def _normalized(value: Decimal) -> Decimal:
    with localcontext() as ctx:
        ctx.prec = 80
        return value.normalize()


def _offset(value: str):
    if value == "Europe/Berlin":
        return ZoneInfo(value)
    if value == "UTC":
        return timezone.utc
    match = re.fullmatch(r"UTC([+-])(\d{2}):(\d{2})", value)
    if not match or int(match[2]) > 14 or int(match[3]) > 59:
        raise ValueError("Unbekannte Zeitbasis.")
    delta = timedelta(hours=int(match[2]), minutes=int(match[3]))
    return timezone(delta if match[1] == "+" else -delta)


def _timestamp(text: str, basis: str | None) -> datetime:
    try:
        value = datetime.fromisoformat(text.replace("Z", "+00:00"))
    except ValueError:
        raise ValueError("Zeitpunkt ist nicht als Datum mit Uhrzeit lesbar.") from None
    if "T" not in text and " " not in text:
        raise ValueError("Ein Datum allein ist kein Transaktionszeitpunkt.")
    if value.tzinfo is not None:
        if basis and value.utcoffset() != value.astimezone(_offset(basis)).utcoffset():
            raise ValueError("Zeitstempel und zugeordnete Zeitbasis widersprechen sich.")
    else:
        if not basis:
            raise ValueError("Zeitbasis fehlt; keine UTC-Annahme fuer Time.")
        zone = _offset(basis)
        first, second = value.replace(tzinfo=zone, fold=0), value.replace(tzinfo=zone, fold=1)
        if first.utcoffset() != second.utcoffset():
            raise ValueError("Mehrdeutiger oder nicht existierender Zeitpunkt beim Zeitwechsel.")
        value = first
    return value.astimezone(timezone.utc)


class MexcNativeImporter:
    def __init__(self, store: DuckDbStore):
        self.store = store
        self.archive = Path(store.path).parent / "mexc_originals"

    def read(self, filename: str, content: bytes, sheet_name: str | None):
        try:
            return self._read(filename, content, sheet_name)
        except (BadZipFile, KeyError, IndexError, StopIteration, ET.ParseError, UnicodeError):
            raise ValueError("MEXC-Datei nicht lesbar; unverschluesseltes XLSX/UTF-8-CSV mit Originalkopfzeile verwenden.") from None

    def _read(self, filename: str, content: bytes, sheet_name: str | None):
        if len(content) > 25 * 1024 * 1024:
            raise ValueError("MEXC-Datei ist groesser als 25 MiB.")
        if Path(filename).suffix.lower() == ".csv":
            text = content.decode("utf-8-sig")
            from steuerreport_v2.importers.csv_raw_importer import CsvRawImporter
            reader = csv.DictReader(StringIO(text), dialect=CsvRawImporter._dialect_for_text(text))
            columns = list(reader.fieldnames or [])
            rows = []
            for row in reader:
                if None in row:
                    raise ValueError("MEXC-CSV enthaelt mehr Werte als Spalten.")
                rows.append((reader.line_num, dict(row)))
                if len(rows) > MAX_ROWS:
                    raise ValueError("MEXC: maximal 10000 Datenzeilen je Datei; Export bitte aufteilen.")
            return rows, columns, "csv", None, []
        if Path(filename).suffix.lower() != ".xlsx":
            raise ValueError("Native MEXC-Dateien bitte als XLSX oder UTF-8-CSV bereitstellen.")
        with ZipFile(BytesIO(content)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError("MEXC-XLSX entpackt groesser als 100 MiB.")
            book = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
            try:
                names = book.sheetnames
                selected = sheet_name or names[0]
                if selected not in names:
                    raise ValueError("Tabellenblatt nicht vorhanden.")
                sheet = book[selected]
                if sheet.max_row and sheet.max_row > MAX_ROWS + 1:
                    raise ValueError("MEXC: maximal 10000 Datenzeilen je Datei; Export bitte aufteilen.")
                if sheet.max_column and sheet.max_column > 256:
                    raise ValueError("MEXC: hoechstens 256 Spalten.")
                # Numerische XML-Literale erhalten; kein Float-Roundtrip fuer Mengen oder UID.
                workbook = ET.fromstring(archive.read("xl/workbook.xml"))
                sheet_id = next(x.attrib["{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"]
                                for x in workbook.findall("s:sheets/s:sheet", NS) if x.attrib["name"] == selected)
                relations = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
                target = next(x.attrib["Target"] for x in relations if x.attrib["Id"] == sheet_id)
                path = posixpath.normpath(target.lstrip("/") if target.startswith("/") else posixpath.join("xl", target))
                if not path.startswith("xl/worksheets/") or ".." in Path(path).parts:
                    raise ValueError("Ungueltiger Tabellenblattpfad.")
                xml = ET.fromstring(archive.read(path))
                numerics = {c.attrib["r"]: c.find("s:v", NS).text for c in xml.findall("s:sheetData/s:row/s:c", NS)
                            if c.attrib.get("t", "n") == "n" and c.find("s:v", NS) is not None}
                iterator = sheet.iter_rows()
                columns = [str(c.value or "").strip() for c in next(iterator, ())]
                rows = []
                for number, cells in enumerate(iterator, start=2):
                    if not any(c.value is not None for c in cells):
                        continue
                    row = {}
                    for column, cell in zip(columns, cells):
                        if not column:
                            continue
                        value = cell.value
                        if value is None:
                            row[column] = None
                            continue
                        if cell.coordinate in numerics and not cell.is_date:
                            value = numerics[cell.coordinate]
                        elif isinstance(value, datetime):
                            value = value.isoformat(sep=" ")
                        row[column] = None if value is None else str(value)
                    rows.append((number, row))
                return rows, columns, "xlsx", selected, names
            finally:
                book.close()

    def _existing(self):
        with self.store.connect() as con:
            rows = con.execute("""select r.raw_object_id,b.filename,r.row_number,r.payload_json
                from raw_objects r join raw_import_batches b on b.batch_id=r.batch_id
                where b.source='mexc' order by b.created_at,r.row_number,r.raw_object_id""").fetchall()
        index = defaultdict(list)
        roots = {}
        for raw_id, filename, number, payload in rows:
            meta = metadata(json.loads(payload))
            if meta.get("version") != 1:
                raise ValueError("Vorhandene MEXC-Altbelege haben keine native Konto-/Exportartbindung; zuerst gesondert pruefen.")
            if meta.get("version") == 1 and meta.get("role") == "book":
                if not {"uid", "account", "signature", "export_kind", "quantity", "event_time"} <= meta.keys():
                    raise ValueError("Vorhandene MEXC-Belegzuordnung ist unvollstaendig; keine unsichere Dublettenpruefung.")
                roots[raw_id] = meta
                index[meta.get("signature")].append(MexcCandidate(raw_object_id=raw_id,
                    filename=filename, row_number=number, export_kind=meta["export_kind"], provider_id=meta.get("provider_id")))
        return index, roots

    def prepare(self, filename: str, content: bytes, sheet_name=None, mexc=None):
        options = MexcImportOptions.model_validate(mexc or {})
        rows, columns, fmt, sheet, names = self.read(filename, content, sheet_name)
        if len(columns) != len(set(columns)) or "_mexc" in columns or None in columns:
            raise ValueError("Doppelte oder reservierte MEXC-Spalten.")
        cols = set(columns)
        time_column = next((c for c in columns if re.fullmatch(r"Creation Time\(UTC[+-]\d{2}:\d{2}\)", c)), None)
        capital_time = next((c for c in columns if re.fullmatch(r"Time\(UTC[+-]\d{2}:\d{2}\)", c)), None)
        detected = None
        if {"UID", "Crypto", "Transaction Type", "Direction", "Quantity"} <= cols and time_column:
            detected = "statement"
        elif {"UID", "Time", "Crypto", "Type", "Quantity", "Status", "Remark"} <= cols:
            detected = "funding_other"
            time_column = "Time"
        elif {"UID", "Date(UTC+8)", "Account", "Crypto", "Total", "USDT Valuation"} <= cols:
            detected = "asset_snapshot"
        elif {"UID", "Futures Trading Pair", "Crypto", "Fund Type", "Fund Flow Type", "Amount"} <= cols and capital_time:
            detected = "futures_capital_flow"
            time_column = capital_time
        problems = []
        if detected is None:
            problems.append("MEXC-Exportlayout nicht unterstuetzt; keine generische Ersatzinterpretation.")
        if not options.account:
            problems.append("MEXC-Konto ausdruecklich zuordnen.")
        compatible = options.export_kind in STATEMENT_KINDS if detected == "statement" else options.export_kind == detected
        if not options.export_kind or not compatible:
            problems.append("Exportart ausdruecklich passend zum Layout waehlen; Dateiname allein genuegt nicht.")
        basis = options.timezone
        if detected == "asset_snapshot":
            basis = "UTC+08:00"
            if options.timezone and options.timezone != basis:
                problems.append("Snapshotdatum ist laut Header UTC+8; abweichende Zeitbasis nicht uebernehmen.")
        if time_column and "(UTC" in time_column:
            basis = time_column[time_column.index("(") + 1:-1]
            if options.timezone and options.timezone != basis:
                problems.append("Zugeordnete Zeitbasis widerspricht dem festen Headeroffset.")
        if detected == "funding_other" and not basis:
            problems.append("Zeitbasis fuer Time fehlt; Exportzeitzone ausdruecklich zuordnen.")
        sha = "sha256:" + hashlib.sha256(content).hexdigest()
        fingerprint = sha + "#mexc=" + _hash({"sheet": sheet, "kind": options.export_kind,
            "account": options.account, "timezone": basis})
        duplicate = self.store.get_import_batch_by_file_hash("mexc", fingerprint)
        report = MexcImportPreview(detected_kind=detected, raw_rows=len(rows),
            disposition="empty" if not rows else "transactions" if options.export_kind in NATIVE_KINDS else "evidence_only",
            blocking_reasons=problems)
        if not rows:
            report.blocking_reasons.append("Leerer Export: nur Kopfzeile, keine Buchungen.")
        if detected is None:
            report.disposition = "unsupported"
        index, roots = self._existing()
        provider_ids = defaultdict(set)
        for root in roots.values():
            if root.get("provider_id"):
                provider_ids[(root["uid"], root["account"], root["provider_id"])].add(root["signature"])
        seen_provider_ids = set()
        prepared = []
        valid_decisions = set()
        for number, row in rows:
            meta = {"version": 1, "export_kind": options.export_kind, "account": options.account,
                    "timezone": basis, "sheet_name": sheet, "file_sha256": sha,
                    "row_number": number, "file_format": fmt, "role": "review"}
            item = MexcRowPreview(row_number=number, status="review", reason="Vorgang fachlich nicht unterstuetzt; nur Rohbeleg, keine Bestandsbuchung.")
            try:
                self._row(row, detected, time_column, basis, options, meta, item)
                if meta["role"] == "book" and not duplicate:
                    provider = meta.get("provider_id")
                    if provider:
                        key = (meta["uid"], meta["account"], provider)
                        if key in seen_provider_ids or provider_ids.get(key, {meta["signature"]}) != {meta["signature"]}:
                            raise ValueError("Wiederholte oder widerspruechliche Provider-Buchungs-ID.")
                        seen_provider_ids.add(key)
                    candidates = index.get(meta["signature"], [])
                    if provider:
                        candidates = [c for c in candidates if not c.provider_id or c.provider_id == provider]
                    item.candidates = candidates
                    decision = options.decisions.get(str(number))
                    if candidates:
                        valid_decisions.add(str(number))
                        if not decision:
                            item.status = "duplicate_candidate"
                            item.reason = "Moeglicher weiterer Beleg einer vorhandenen Buchung; verknuepfen oder eigenstaendige Buchung begruenden."
                            report.blocking_reasons.append(f"Zeile {number}: moeglicher Doppelbeleg noch nicht zugeordnet.")
                        elif decision.action == "link_existing":
                            if decision.target_raw_object_id not in {x.raw_object_id for x in candidates}:
                                raise ValueError("Linkziel passt nicht zu Konto, UID, Zeitpunkt, Asset, Menge und Quelltyp.")
                            meta.update(role="linked", link_target_raw_object_id=decision.target_raw_object_id,
                                        decision_reason=decision.reason)
                            item.status, item.reason = "linked", "Zusaetzlicher Originalbeleg; keine zweite Bestandsbuchung."
                        else:
                            if decision.target_raw_object_id:
                                raise ValueError("Eigenstaendige Buchung darf kein Linkziel enthalten.")
                            meta["decision_reason"] = decision.reason
                    elif decision:
                        raise ValueError("Zuordnungsentscheidung passt nicht mehr zur aktuellen Vorschau.")
            except (ValueError, InvalidOperation, TypeError):
                meta["role"] = "review"
                item.status = "invalid"
                item.reason = "Pflichtwerte, Richtung, Zeitpunkt oder Belegzuordnung ungueltig; Original und Zuordnung pruefen."
                report.blocking_reasons.append(f"Zeile {number}: {item.reason}")
            report.rows.append(item)
            report.bookable_rows += meta["role"] == "book" and item.status != "duplicate_candidate"
            report.linked_rows += meta["role"] == "linked"
            report.snapshot_rows += meta["role"] == "snapshot"
            report.bonus_rows += meta["role"] == "bonus"
            report.review_rows += meta["role"] in {"book", "bonus", "review"}
            prepared.append((number, row, meta))
        if not duplicate and set(options.decisions) - valid_decisions:
            report.blocking_reasons.append("Zeilenentscheidungen sind veraltet oder beziehen sich nicht auf passende Kandidaten.")
        report.warnings = ["Mengenabgleich ist keine Kostenbasis oder Steuerfreigabe. Airdrop bleibt ungeklart; keine automatische Mining-/Reward-Zuordnung.",
            "Self-Service-Export laut MEXC ab 01.10.2024; angeforderter Zeitraum beweist keine vollstaendige Historie."]
        preview = CsvImportPreview(source="mexc", filename=filename, profile_label="MEXC Account Data Export",
            file_format=fmt, sheet_name=sheet, sheet_names=names, row_count=len(rows), columns=columns,
            matched_columns={"time": time_column, "asset": "Crypto" if "Crypto" in cols else None,
                "quantity": "Quantity" if "Quantity" in cols else "Amount" if "Amount" in cols else None},
            missing_required_groups=[] if detected else ["mexc_layout"], sample_rows=[r for _, r in rows[:5]],
            can_import=not report.blocking_reasons and bool(rows), duplicate_batch_id=duplicate.batch_id if duplicate else None,
            file_sha256=sha, import_fingerprint=fingerprint, mexc=report)
        return preview, prepared

    @staticmethod
    def _row(row, detected, time_column, basis, options, meta, item):
        uid, asset = str(row.get("UID") or "").strip(), str(row.get("Crypto") or "").strip()
        if not uid or len(uid) > 128 or not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._:-]{0,63}", asset):
            raise ValueError("UID/Asset fehlt.")
        meta.update(uid=uid, asset=asset.upper())
        provider_id = str(row.get("Transaction ID") or row.get("Record ID") or "").strip()
        if len(provider_id) > 256:
            raise ValueError("Buchungs-ID zu lang.")
        if provider_id:
            meta["provider_id"] = provider_id
        item.asset = asset.upper()
        if detected == "asset_snapshot":
            if not re.fullmatch(r"\d{8}", str(row.get("Date(UTC+8)") or "")):
                raise ValueError("Snapshotdatum muss YYYYMMDD entsprechen.")
            datetime.strptime(str(row.get("Date(UTC+8)") or ""), "%Y%m%d")
            meta.update(role="snapshot", snapshot_date=row["Date(UTC+8)"])
            item.status, item.reason = "evidence_only", "Bestandsnachweis; Total und USDT Valuation erzeugen keine Erwerbe oder Kostenbasis."
            return
        if not time_column:
            return
        instant = _timestamp(str(row.get(time_column) or ""), basis)
        text = str(row.get("Quantity") if "Quantity" in row else row.get("Amount") or "").strip()
        if not re.fullmatch(r"[+-]?\d+(?:\.\d+)?(?:[eE][+-]?\d+)?", text):
            raise ValueError("Menge ist keine eindeutige Dezimalzahl.")
        quantity = Decimal(text)
        if not quantity.is_finite() or quantity.copy_abs() >= Decimal("1e20") or _normalized(quantity).as_tuple().exponent < -18:
            raise ValueError("Menge ausserhalb verlustfreier Datenbankpraezision.")
        direction = str(row.get("Direction") or row.get("Fund Flow Type") or "").strip().lower()
        if direction and direction not in {"inflow", "outflow"}:
            raise ValueError("Richtung unbekannt.")
        if detected == "statement" and not direction:
            raise ValueError("Richtung fehlt.")
        if direction == "inflow" and quantity < 0:
            raise ValueError("Negativer Inflow widerspricht der Richtung.")
        if direction == "outflow":
            quantity = quantity.copy_abs().copy_negate()
        label = str(row.get("Transaction Type") or row.get("Type") or row.get("Fund Type") or "").strip()
        meta.update(event_time=instant.isoformat(), quantity=str(quantity), label=label)
        item.quantity, item.event_time = str(quantity), instant.isoformat()
        if detected == "futures_capital_flow" and label.upper() == "BONUS":
            meta["role"] = "bonus"
            item.status, item.reason = "evidence_only", "BONUS-Zuteilung/Ruecknahme: Rechte und Kontowirkung offen; keine Gewinne oder Lots."
            return
        if options.export_kind not in NATIVE_KINDS or label.lower() != "airdrop" or quantity == 0:
            return
        meta["role"] = "book"
        meta["signature"] = _hash([uid, options.account, instant.isoformat(), asset.upper(), str(_normalized(quantity)), label.lower()])
        item.status, item.reason = "ready", "Belegte Mengenbewegung, steuerliche Ursache offen; unknown, kein Mining/Reward."

    def import_file(self, filename: str, content: bytes, sheet_name=None, mexc=None):
        with LOCK:
            preview, prepared = self.prepare(filename, content, sheet_name, mexc)
            if preview.duplicate_batch_id:
                return self.store.get_import_batch(preview.duplicate_batch_id)
            if not preview.can_import:
                raise ValueError("; ".join(preview.mexc.blocking_reasons))
            self.archive.mkdir(parents=True, exist_ok=True, mode=0o700)
            original = self.archive / (preview.file_sha256.removeprefix("sha256:") + ".original")
            try:
                fd = os.open(original, os.O_WRONLY | os.O_CREAT | os.O_EXCL | os.O_NOFOLLOW, 0o600)
            except FileExistsError:
                if original.is_symlink() or hashlib.sha256(original.read_bytes()).hexdigest() != preview.file_sha256[7:]:
                    raise ValueError("Gespeicherter MEXC-Originalbeleg stimmt nicht mit seinem Hash ueberein.") from None
            else:
                with os.fdopen(fd, "wb") as handle:
                    handle.write(content)
                    handle.flush()
                    os.fsync(handle.fileno())
            return self.store.create_import_batch_with_raw(ImportBatchCreate(source="mexc", filename=filename, file_hash=preview.import_fingerprint), [RawObjectCreate(row_number=number,
                payload_json=json.dumps({**row, "_mexc": json.dumps(meta, sort_keys=True)}, ensure_ascii=False, sort_keys=True))
                for number, row, meta in prepared])


def event_from_raw(batch, raw) -> CanonicalEventCreate | None:
    row = json.loads(raw.payload_json)
    meta = metadata(row)
    if meta.get("version") != 1 or meta.get("role") != "book":
        return None
    if meta.get("row_number") != raw.row_number or not str(batch.file_hash or "").startswith(str(meta.get("file_sha256")) + "#mexc="):
        return None
    expected_context = _hash({"sheet": meta.get("sheet_name"), "kind": meta.get("export_kind"),
        "account": meta.get("account"), "timezone": meta.get("timezone")})
    if batch.file_hash != meta["file_sha256"] + "#mexc=" + expected_context:
        return None
    try:
        options = MexcImportOptions(export_kind=meta["export_kind"], account=meta["account"])
        time_column = "Time" if options.export_kind == "funding_other" else next(
            key for key in row if re.fullmatch(r"Creation Time\(UTC[+-]\d{2}:\d{2}\)", key))
        check = {}
        MexcNativeImporter._row(row, "funding_other" if options.export_kind == "funding_other" else "statement",
            time_column, meta["timezone"], options, check, MexcRowPreview(row_number=raw.row_number, status="review", reason=""))
        if any(check.get(key) != meta.get(key) for key in ("role", "uid", "asset", "quantity", "event_time", "label", "signature", "provider_id")):
            return None
    except (ValueError, KeyError, StopIteration, TypeError, InvalidOperation):
        return None
    return CanonicalEventCreate(source="mexc", source_ref="mexc:v1:" + meta["signature"] + ":" + str(raw.row_number),
        event_time=datetime.fromisoformat(meta["event_time"]), event_type="unknown", asset=meta["asset"],
        quantity=Decimal(meta["quantity"]), account="mexc:" + quote(meta["uid"], safe="") + ":" + quote(meta["account"], safe=""), raw_object_id=raw.raw_object_id)


def evidence_for_batch(store, batch_id):
    with store.connect() as con:
        rows = con.execute("""select r.raw_object_id,r.batch_id,r.row_number,r.payload_json,b.filename,b.file_hash,e.event_id
            from raw_objects r join raw_import_batches b on b.batch_id=r.batch_id
            left join canonical_events e on e.raw_object_id=r.raw_object_id where b.source='mexc'
            order by b.created_at,r.row_number""").fetchall()
    records = []
    for raw_id, raw_batch, number, payload, filename, fingerprint, event_id in rows:
        meta = metadata(json.loads(payload))
        records.append({"raw_object_id": raw_id, "batch_id": raw_batch, "row_number": number,
            "filename": filename, "import_fingerprint": fingerprint, "event_id": event_id, "metadata": meta})
    own_ids = {r["raw_object_id"] for r in records if r["batch_id"] == batch_id}
    linked_ids = {r["metadata"].get("link_target_raw_object_id") for r in records if r["batch_id"] == batch_id}
    return [r for r in records if r["batch_id"] == batch_id or r["raw_object_id"] in linked_ids
            or r["metadata"].get("link_target_raw_object_id") in own_ids]
